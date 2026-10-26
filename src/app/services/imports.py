"""Import a list of emerging technologies: an Excel file laid out like Tata's
"New Emerging Technologies" sheet — Technology Name, then Domain,
Sub-domain, Description, TRL and Commercial Availability.

  upload    the sheet whose header row has "Technology Name" is read; any
            other sheet (such as Technologies Tracked) is ignored, and a file
            that doesn't fit is refused with the reason
  check     names trimmed, blank rows dropped, repeats merged
  match     an exact name or alias costs nothing, and so does a proposal from
            the name outside or inside its brackets ("DNSSEC"); the rest go
            to the organize skill against the whole register, which proposes
            matches. Code keeps a narrower name ("Telco edge computing",
            "Computer vision for robotics") apart from what it proposes.
  review    nothing is saved yet: proposed matches are accepted or rejected
            (a rejected one comes in as new), and values not in our lists are
            mapped; a sub-domain not in our lists waits for a person's choice
  confirm   a match takes the file's wording as an alias; a new one is added
            with no pages. A row that alone names its technology has its
            values recorded as a person's choices (they win over what the
            model judges); where several rows name one technology, their
            values stay on this import's record.
  reopen    undoes what confirm did for rows that joined a technology and
            puts them back to review, to confirm again

An import means "these are emerging technologies Tata cares about": each
row's own commercial availability is kept. Importing again updates and never
duplicates, and never undoes a person's later change to an imported value."""

import io
import json
import re
import uuid
from collections import defaultdict

import openpyxl
from pymongo import ReplaceOne

from app.services import auth, classify_effective, decisions, jobs
from app.services.classify_effective import IMPORTED
from app.services.db.mongo import get_collection
from app.services.timezone import now_utc
from app.services.display import display_name

HEADERS = ["Technology Name", "Domain", "Sub-domain", "Description", "TRL", "Commercial Availability"]
FIELD = {"Domain": "domain", "Sub-domain": "subdomain", "Description": "description", "TRL": "trl",
         "Commercial Availability": "availability"}
ALIASES = {"blockchain": "blockchain_web3", "blockchain/cybersecurity": "blockchain_web3", "blockchain/web3": "blockchain_web3"}
QUALIFIERS = {"for", "in", "at", "on"}  # "X for Y": what follows says where X is used, so it's narrower
CHARS_PER_TOKEN = 4                     # the rule of thumb the other cost previews use
ADDED_BY_HAND = "Added by hand"


class Refused(ValueError):
    """A file or an action that can't be done; the message says why."""


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def read(content: bytes, filename: str) -> dict:
    """{sheet, rows: [{row, name, values}], columns}."""
    if not filename.lower().endswith(".xlsx"):
        raise Refused("only .xlsx files can be imported, laid out like Tata's New Emerging Technologies sheet")
    try:
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise Refused("the file couldn't be opened as an Excel workbook")
    for ws in book.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        head = [str(h or "").strip().lower() for h in rows[0]]
        if "technology name" not in head:
            continue
        cols = {h: head.index(h.lower()) for h in HEADERS if h.lower() in head}
        out = []
        for n, r in enumerate(rows[1:], start=2):
            cell = lambda h: str(r[cols[h]]).strip() if h in cols and cols[h] < len(r) and r[cols[h]] is not None else ""
            name = cell("Technology Name")
            if name:
                out.append({"row": n, "name": name, "values": {FIELD[h]: cell(h) for h in HEADERS[1:] if cell(h)}})
        return {"sheet": ws.title, "rows": out, "columns": [h for h in HEADERS if h in cols]}
    raise Refused("no 'Technology Name' column in any sheet")


def clean(rows: list[dict]) -> list[dict]:
    """Repeats (the same name, any case) merged: the first row's values, filled
    in from later ones."""
    seen: dict[str, dict] = {}
    for r in rows:
        key = r["name"].lower()
        if key in seen:
            seen[key]["values"] = {**r["values"], **seen[key]["values"]}
            seen[key].setdefault("repeats", []).append(r["row"])
        else:
            seen[key] = {**r, "name": re.sub(r"\s+", " ", r["name"]).strip()}
    return list(seen.values())


def register_index() -> dict[str, dict]:
    idx = {}
    for t in get_collection("technologies").find({}, {"name": 1, "aliases": 1, "label": 1}):
        for n in {t["name"], *(t.get("aliases") or []), (t.get("label") or "").lower()}:
            if n:
                idx.setdefault(n.lower().strip(), t)
    return idx


# ── matching: what costs nothing, and code's check on the skill ──────

def _words(text: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", text.lower())


def _forms(name: str) -> list[list[str]]:
    """A name's words outside its brackets, then the words in each bracket
    on their own (often its abbreviation)."""
    outer = _words(re.sub(r"\([^)]*\)", " ", name))
    return [f for f in (outer, *map(_words, re.findall(r"\(([^)]*)\)", name))) if f]


def narrower(name: str, names: list[str]) -> bool:
    """Code's check on a match the organize skill proposes: the file's name is
    one of the technology's names with a word in front ("Telco edge
    computing" → "edge computing"), or says what it is for, in, at or on that
    none of the technology's names do ("Computer vision for robotics", "AI in
    education"). Either way it is a narrower technology, never the same one."""
    mine, theirs = _forms(name), [f for n in names if n for f in _forms(n)]
    if not mine or any(f == t for f in mine for t in theirs):
        return False
    outer = mine[0]
    if any(len(t) < len(outer) and outer[-len(t):] == t for t in theirs):
        return True
    q = next((i for i, w in enumerate(outer) if i and w in QUALIFIERS), None)
    known = {w for t in theirs for w in t}
    return q is not None and not any(w in known for w in outer[q + 1:])


def free_match(name: str, idx: dict[str, dict] | None = None) -> tuple[str, dict | None, str | None]:
    """(group, technology, why) for one name, at no cost: "matched" when it is
    a technology's name or alias; "proposed" when its name outside its
    brackets, or the name inside them, is (a person confirms it); else
    "unmatched", for the organize skill."""
    idx = register_index() if idx is None else idx
    hit = idx.get(name.lower().strip())
    if hit:
        return "matched", hit, None
    outer = re.sub(r"\s+", " ", re.sub(r"\([^)]*\)", " ", name)).strip()
    inner = [i.strip() for i in re.findall(r"\(([^)]*)\)", name)]
    for form, why in ((outer, "Without the words in brackets, it is “{}”, a technology in the register"),
                      *((i, "The name in its brackets, “{}”, is a technology in the register") for i in inner)):
        hit = idx.get(form.lower()) if form and form.lower() != name.lower() else None
        if hit:
            return "proposed", hit, why.format(form)
    return "unmatched", None, None


def match_cost(names: list[str]) -> dict:
    """What matching these names with the organize skill costs, before it runs:
    the calls, and about how many tokens they send (the skill's text and the
    whole register in each call, plus the names)."""
    from app.services.prompts.store import PromptStore
    from app.settings import get_settings
    calls = -(-len(names) // get_settings().resolve_batch_size) if names else 0
    register = [t["name"] for t in get_collection("technologies").find({}, {"name": 1})]
    try:
        p = PromptStore.get("technology_resolve")
        skill = len(p.system or "") + len(p.prompt)
    except KeyError:
        skill = 0
    chars = calls * (skill + len(json.dumps(register, ensure_ascii=False))) + len(json.dumps(names, ensure_ascii=False))
    return {"names": len(names), "calls": calls, "register": len(register), "tokens_about": chars // CHARS_PER_TOKEN}


def propose(names: list[str], run_id: str, llm=None) -> tuple[dict[str, dict], dict[str, dict], int]:
    """({name: the register technology the organize skill folds it into},
    {name: a technology it was folded into but code kept it apart from, as
    narrower}, calls made): lowercase names, in batches; a name it folds
    nowhere is left out of both."""
    from agent.resolver.resolver import _chunks, _llm_groups
    from app.services.llm.context import with_context
    from app.services.llm.llm import LLM
    from app.services.prompts.store import PromptStore
    from app.settings import get_settings
    register = {t["name"]: t for t in get_collection("technologies").find({}, {"name": 1, "label": 1, "aliases": 1})}
    llm = llm or LLM.initialize()
    prompt = PromptStore.get("technology_resolve")
    found, kept, calls = {}, {}, 0
    with with_context(run_kind="import", run_id=run_id, skills=["technology_resolve"]):
        for chunk in _chunks(sorted(names), get_settings().resolve_batch_size):
            calls += 1
            for g in _llm_groups(llm, prompt, chunk, sorted(register)):
                target = register.get(g["name"])
                for m in g["members"]:
                    if target and m != g["name"]:
                        own = [target["name"], target.get("label") or "", *(target.get("aliases") or [])]
                        (kept if narrower(m, own) else found)[m] = target
    return found, kept, calls


def target_info(tech_ids: list[str]) -> dict[str, dict]:
    """What each technology a row would join is now, to judge a match by: its
    pages, and its classification with where each answer came from."""
    techs = list(get_collection("technologies").find({"_id": {"$in": tech_ids}}, {
        "name": 1, "label": 1, "page_count": 1, "origin": 1, "classification": 1, "assessed_classification": 1}))
    decs, lists = decisions.active([t["_id"] for t in techs]), classify_effective.names()
    out = {}
    for t in techs:
        e = classify_effective.effective(t, decs, lists)
        out[t["_id"]] = {"name": display_name(t), "page_count": t.get("page_count") or 0,
                         "origin": t.get("origin"),
                         "availability": {"name": e["availability"]["name"], "source": e["availability"]["source"]},
                         "domain": {"name": e["domain"]["name"], "subdomain": e["domain"]["subdomain"],
                                    "source": e["domain"]["source"]},
                         "trl": {"level": e["trl"]["level"], "source": e["trl"]["source"]}}
    return out


# ── values: map the file's words to our lists ────────────────────────

def _lists() -> dict:
    return {"availability": {a["name"].lower(): a["_id"] for a in get_collection("availability").find()} |
            {a["_id"]: a["_id"] for a in get_collection("availability").find()},
            "domain": {d["name"].lower(): d["_id"] for d in get_collection("domains").find()} | ALIASES,
            "subdomain": {(s["domain_id"], s["name"].lower()): s["_id"] for s in get_collection("subdomains").find()}}


def _key(u: dict) -> tuple:
    return u["field"], u["value"], u.get("domain_id") if u["field"] == "subdomain" else None


def map_values(items: list[dict], chosen: list[dict] | None = None) -> list[dict]:
    """Each value the file uses that our lists don't have, with the choice made
    so far (`chosen`: the review's list). A sub-domain is looked at under its
    domain, so a domain mapped in the review brings its sub-domains in; a
    sub-domain not in our lists waits for a person (keep None) until they
    map it, keep it as written, or leave it out."""
    ls = _lists()
    before = {_key(u): u for u in chosen or []}
    domain_to = {u["value"]: u["map_to"] for u in chosen or [] if u["field"] == "domain" and u.get("map_to")}
    unknown: dict[tuple, dict] = {}

    def add(field: str, value: str, **extra) -> None:
        u = unknown.setdefault((field, value, extra.get("domain_id")),
                               {"field": field, "value": value, **extra, "count": 0, "map_to": None})
        u["count"] += 1

    for it in items:
        v = it["values"]
        if v.get("availability") and v["availability"].lower() not in ls["availability"]:
            add("availability", v["availability"])
        dom = ls["domain"].get((v.get("domain") or "").lower()) or domain_to.get(v.get("domain") or "")
        if v.get("domain") and not ls["domain"].get(v["domain"].lower()):
            add("domain", v["domain"])
        if v.get("subdomain") and dom and (dom, v["subdomain"].lower()) not in ls["subdomain"]:
            add("subdomain", v["subdomain"], domain_id=dom, keep=None)
        if v.get("trl") and not (v["trl"].isdigit() and 1 <= int(v["trl"]) <= 9):
            add("trl", v["trl"])
    for k, u in unknown.items():
        if k in before:
            u["map_to"], u["keep"] = before[k].get("map_to"), before[k].get("keep", u.get("keep"))
    return list(unknown.values())


def resolved_values(values: dict, unknown: list[dict]) -> dict:
    """The row's values as list ids, using the review's mapping."""
    ls = _lists()
    mapped = {_key(u): u for u in unknown}
    out = {}
    av = values.get("availability")
    if av:
        out["availability"] = ls["availability"].get(av.lower()) or (mapped.get(("availability", av, None)) or {}).get("map_to")
    dom = values.get("domain")
    if dom:
        out["domain"] = ls["domain"].get(dom.lower()) or (mapped.get(("domain", dom, None)) or {}).get("map_to")
    sub = values.get("subdomain")
    if sub and out.get("domain"):
        u = mapped.get(("subdomain", sub, out["domain"])) or {}
        out["subdomain"] = ls["subdomain"].get((out["domain"], sub.lower())) or u.get("map_to") or \
            (f"new:{sub}" if u.get("keep") else None)
    trl = values.get("trl")
    if trl:
        out["trl"] = int(trl) if trl.isdigit() and 1 <= int(trl) <= 9 else (mapped.get(("trl", trl, None)) or {}).get("map_to")
    return {k: v for k, v in out.items() if v not in (None, "")}


def chosen_values(availability: str | None = None, domain: str | None = None, subdomain: str | None = None,
                  trl: int | None = None) -> dict:
    """Values picked from our lists (the add-by-hand form), checked: each must
    be in its list, and a sub-domain under the domain chosen."""
    ls = _lists()
    if availability and availability not in ls["availability"].values():
        raise Refused(f"{availability!r} isn't a commercial availability in our list")
    if domain and domain not in ls["domain"].values():
        raise Refused(f"{domain!r} isn't a domain in our list")
    if subdomain and not domain:
        raise Refused("choose the domain a sub-domain belongs to")
    if subdomain and (domain, subdomain) not in {(d, s) for (d, _), s in ls["subdomain"].items()}:
        raise Refused(f"{subdomain!r} isn't a sub-domain of {domain!r} in our list")
    if trl is not None and not 1 <= trl <= 9:
        raise Refused("TRL is a whole number from 1 to 9")
    out = {"availability": availability, "domain": domain, "subdomain": subdomain, "trl": trl}
    return {k: v for k, v in out.items() if v not in (None, "")}


def _changed_since(tech_id: str, field: str, value, active: dict | None) -> bool:
    """A person undid this value after an import recorded it, or chose another
    since: importing it again doesn't undo their choice."""
    prior = list(get_collection("decisions").find(
        {"kind": "move", "tech_id": tech_id, "field": field, "value": value, "reason": {"$regex": f"^{re.escape(IMPORTED)}"},
         "reopened_at": None}, {"at": 1, "undone_at": 1}))
    if any(d.get("undone_at") for d in prior):
        return True
    return bool(active) and classify_effective.imported_from(active) is None and any(d["at"] < active["at"] for d in prior)


def set_values(tech_id: str, values: dict, reason: str, tz: str, extra: dict | None = None,
               imported: bool = False) -> int:
    """Record values as a person's choices where they differ from the ones in
    force; returns how many. An imported value (`imported`) isn't recorded
    again where a person has changed or undone it since."""
    active = {d["field"]: d for d in decisions.active([tech_id], ["move"])}
    n = 0
    for field, value in values.items():
        a = active.get(field)
        if (a and a.get("value") == value) or (imported and _changed_since(tech_id, field, value, a)):
            continue
        try:
            decisions.create("move", tech_id, field=field, value=value, reason=reason, tz=tz, extra=extra)
            n += 1
        except decisions.Refused:
            pass  # a value our lists can't hold is kept on the import record only
    return n


# ── the import record ─────────────────────────────────────────────────

ITEMS = "import_items"  # an import's rows, each its own record (from 24 Oct): no import outgrows a record


def load(import_id: str) -> dict | None:
    """An import with its rows (`items`), in row order. Imports saved before
    24 Oct hold their rows inside; they move out the next time they're saved."""
    imp = get_collection("imports").find_one({"_id": import_id})
    if imp is not None and "items" not in imp:
        imp["items"] = list(get_collection(ITEMS).find({"import_id": import_id}, {"_id": 0, "import_id": 0}).sort("row", 1))
    return imp


def save(imp: dict, rows: set[int] | None = None) -> None:
    """Write an import back: its rows apart (only `rows` when given: the ones
    that changed), the rest as one small record."""
    import_id = imp["_id"]
    moving = "items" in (get_collection("imports").find_one({"_id": import_id}, {"items": {"$slice": 1}}) or {})
    ops = [ReplaceOne({"_id": f"{import_id}:{it['row']}"}, {**it, "_id": f"{import_id}:{it['row']}", "import_id": import_id},
                      upsert=True)
           for it in imp["items"] if rows is None or moving or it["row"] in rows]
    for i in range(0, len(ops), 1000):
        get_collection(ITEMS).bulk_write(ops[i:i + 1000], ordered=False)
    get_collection("imports").replace_one({"_id": import_id}, {k: v for k, v in imp.items() if k not in ("items", "targets")},
                                          upsert=True)


def create(content: bytes, filename: str, tz: str) -> dict:
    data = read(content, filename)
    rows = clean(data["rows"])
    idx = register_index()
    items = []
    for r in rows:
        group, hit, why = free_match(r["name"], idx)
        items.append({**r, "group": group, "target_tech_id": hit["_id"] if hit else None,
                      "target_name": (hit or {}).get("label") or (hit or {}).get("name"), "why": why,
                      "decision": None})
    doc = {"_id": uuid.uuid4().hex[:12], "file": filename, "sheet": data["sheet"], "columns": data["columns"], "at": now_utc(),
           "tz": tz, "by": auth.by(), "status": "reviewing", "rows": len(data["rows"]), "items": items,
           "unknown_values": map_values(items), "calls": 0}
    doc["counts"] = counts(doc)
    save(doc)
    return doc


def counts(imp: dict) -> dict:
    out = {"matched": 0, "proposed": 0, "new": 0, "unmatched": 0, "accepted": 0, "rejected": 0}
    for it in imp["items"]:
        out[it["group"]] += 1
        if it["decision"] in ("accept", "reject"):
            out[it["decision"] + "ed"] += 1
    waiting = [u for u in imp["unknown_values"] if u["field"] == "subdomain" and u.get("keep") is None and not u.get("map_to")]
    out["values_waiting"] = len(waiting)
    out["values_to_map"] = sum(1 for u in imp["unknown_values"] if not u.get("map_to") and not u.get("keep")) - len(waiting)
    return out


def view(imp: dict) -> dict:
    """The import as the review screen shows it: with what each technology a
    row would join is now, and the newer import that replaced it, if any."""
    return {**imp, "targets": target_info(sorted({it["target_tech_id"] for it in imp["items"] if it.get("target_tech_id")})),
            "replaced_by": replaced_by(imp)}


def waiting(imp: dict) -> int:
    """How many proposed matches wait for a person's decision (none once it is
    confirmed or discarded)."""
    if imp.get("status") not in ("reviewing", "ready"):
        return 0
    if "items" not in imp:  # its rows are apart: count them there
        return get_collection(ITEMS).count_documents({"import_id": imp["_id"], "group": "proposed", "decision": None})
    return sum(1 for i in imp["items"] if i.get("group") == "proposed" and i.get("decision") is None)


def replaced_by(imp: dict) -> dict | None:
    """An import still open for review is replaced when a newer import of the
    same file has been confirmed: that one settled every row this one asks
    about. Its id and when it was confirmed, or None."""
    if imp.get("status") not in ("reviewing", "ready"):
        return None
    newer = get_collection("imports").find_one(
        {"file": imp.get("file"), "status": "confirmed", "at": {"$gt": imp["at"]}},
        {"at": 1, "tz": 1, "confirmed_at": 1, "confirmed_tz": 1}, sort=[("at", -1)])
    return {"id": newer["_id"], "at": newer["at"], "tz": newer.get("tz"), "confirmed_at": newer.get("confirmed_at"),
            "confirmed_tz": newer.get("confirmed_tz")} if newer else None


def new_technology(name: str, origin: str, extra: dict) -> str:
    """A technology added with no pages, its id from its name; returns the id."""
    techs = get_collection("technologies")
    tech_id = slug(name)
    while techs.find_one({"_id": tech_id}):
        tech_id += "-x"
    techs.insert_one({"_id": tech_id, "name": name.lower(), "label": name, "aliases": [name.lower()], "enabled": True,
                      "origin": origin, "page_count": 0, "source_count": 0, "doc_count": 0, "created_at": now_utc(),
                      **extra})
    return tech_id


def match(import_id: str, llm=None) -> dict:
    """Proposed matches for the names without an exact one: the organize skill
    folds file names into the register; a name it folds nowhere, or that code
    keeps apart as narrower, is new."""
    imp = load(import_id)
    names = {it["name"].lower(): it for it in imp["items"] if it["group"] == "unmatched"}
    found, kept, calls = propose(list(names), import_id, llm)
    for name, it in names.items():
        target = found.get(name)
        if target:
            it.update(group="proposed", target_tech_id=target["_id"], target_name=display_name(target))
        else:
            it["group"] = "new"
            if name in kept:
                it["kept_apart"] = {"tech_id": kept[name]["_id"], "name": kept[name].get("label") or kept[name]["name"]}
    imp["calls"] = imp.get("calls", 0) + calls
    imp["status"] = "ready"
    imp["counts"] = counts(imp)
    save(imp, {it["row"] for it in names.values()})
    return imp


def _joins(it: dict) -> bool:
    return it["group"] == "matched" or (it["group"] == "proposed" and it["decision"] == "accept")


def _joining(imp: dict) -> dict[str, list[dict]]:
    """{technology id: the rows that join it}, in file order."""
    out: dict[str, list[dict]] = defaultdict(list)
    for it in imp["items"]:
        if _joins(it):
            out[it["target_tech_id"]].append(it)
    return out


def _record(imp: dict, its: list[dict], at: str, tz: str | None) -> dict:
    """What a technology keeps of the import that names it: the file and the
    rows' names, and the row's values when one row alone names it."""
    rec = {"import_id": imp["_id"], "file": imp["file"], "at": at, "tz": tz, "rows": [it["name"] for it in its]}
    if len(its) == 1:
        rec["row"] = {**its[0]["values"], "name": its[0]["name"]}
    return rec


def confirm(import_id: str, tz: str) -> dict:
    """Apply it: aliases, new technologies, and each row's values as a
    person's choices — when the row alone names its technology. Where several
    rows name one technology, which row's values are right is a person's
    call, so they stay on this import's record. Values already in force, or
    changed by a person since an import recorded them, aren't recorded again,
    so confirming the same rows again changes nothing."""
    from agent.resolver import restamp
    imp = load(import_id)
    techs = get_collection("technologies")
    at, joining = now_utc(), _joining(imp)
    made = joined = shared = 0
    for it in imp["items"]:
        if _joins(it):
            tech_id, its = it["target_tech_id"], joining[it["target_tech_id"]]
            techs.update_one({"_id": tech_id}, {"$addToSet": {"aliases": it["name"].lower()},
                                                "$set": {"imported": _record(imp, its, at, tz), "updated_at": at}})
            joined += 1
        else:
            its, name = [it], it["name"].lower()
            existing = techs.find_one({"$or": [{"name": name}, {"aliases": name}]})
            if existing:
                tech_id = existing["_id"]
                techs.update_one({"_id": tech_id}, {"$set": {"imported": _record(imp, its, at, tz)}})
            else:
                tech_id = new_technology(it["name"], "imported", {"imported": _record(imp, its, at, tz)})
                made += 1
        if len(its) > 1:
            shared += 1
            continue
        set_values(tech_id, resolved_values(it["values"], imp["unknown_values"]), f"{IMPORTED}{imp['file']}", tz,
                   extra={"import_id": import_id, "row": it["row"]}, imported=True)
    result = {"joined": joined, "added": made, "shared": shared}
    get_collection("imports").update_one({"_id": import_id}, {"$set": {"status": "confirmed", "confirmed_at": at,
                                                                     "confirmed_tz": tz, "confirmed_by": auth.by(),
                                                                     "result": result}})
    jobs.later(restamp, [it["name"].lower() for it in imp["items"]])  # the pages that use the rows' names
    return result


# ── reopen: undo what confirm did for some rows, to review them again ──

def _import_decisions(imp: dict) -> list[dict]:
    """The values this import's confirm recorded, oldest first: each says so,
    or (recorded before they did) by its reason and time."""
    return list(get_collection("decisions").find({"kind": "move", "$or": [
        {"extra.import_id": imp["_id"]},
        {"extra": None, "reason": f"{IMPORTED}{imp['file']}", "at": {"$gte": imp["at"], "$lte": imp["confirmed_at"]}},
    ]}).sort("at", 1))


def _made_by(imp: dict, decs: list[dict]) -> dict[str, int]:
    """{decision id: the row whose value it recorded}. A decision says so; one
    recorded before decisions did is traced: confirm went through the rows in
    file order, so it goes to the first row, from the last one traced on that
    technology on, that has its value."""
    joining, last, out = _joining(imp), defaultdict(int), {}
    for d in decs:
        row = (d.get("extra") or {}).get("row")
        if row is None:
            its = joining.get(d["tech_id"], [])
            i = next((i for i in range(last[d["tech_id"]], len(its))
                      if resolved_values(its[i]["values"], imp["unknown_values"]).get(d["field"]) == d["value"]), None)
            if i is None:
                continue
            last[d["tech_id"]], row = i, its[i]["row"]
        out[d["_id"]] = row
    return out


def _restore_own(d: dict, mine: set[str]) -> bool:
    """Bring back the choice that was in force on this field before this
    import recorded its first value there (a person's, or an earlier
    import's); True if there was one."""
    coll = get_collection("decisions")
    first = coll.find_one({"_id": {"$in": list(mine)}, "tech_id": d["tech_id"], "field": d["field"], "kind": "move"},
                          sort=[("at", 1)])
    own = coll.find_one({"kind": "move", "tech_id": d["tech_id"], "field": d["field"], "fabric_id": None,
                         "at": {"$lt": first["at"]}, "_id": {"$nin": list(mine)}, "reopened_at": None},
                        sort=[("at", -1)])
    if own is None or own.get("undone_at"):
        return False
    coll.update_one({"_id": own["_id"]}, {"$set": {"active": True}, "$unset": {"replaced_at": ""}})
    return True


def reopen(import_id: str, rows: list[int] | None, tz: str) -> dict:
    """Undo what confirm did for these rows, and put them back to review:
    the name a proposed match added as an alias is taken off again, the
    values the row recorded on the technology are withdrawn (the technology's
    own come back), and the technology's import record no longer names it.
    `rows`: rows that joined a technology; by default every accepted proposed
    match. An exact match can be named too: only its values are withdrawn
    (its name was an alias already), and confirming again applies it under
    the rule for several rows. Confirming again changes nothing for the
    other rows."""
    from agent.resolver import restamp
    coll, techs = get_collection("decisions"), get_collection("technologies")
    imp = load(import_id)
    if imp is None:
        raise Refused("no such import")
    if imp["status"] != "confirmed":
        raise Refused("only a confirmed import can be reopened: this one hasn't been confirmed")
    by_row = {it["row"]: it for it in imp["items"]}
    rows = rows or [it["row"] for it in imp["items"] if it["group"] == "proposed" and it["decision"] == "accept"]
    if not rows:
        raise Refused("this import has no accepted proposed matches to reopen")
    for r in rows:
        if r not in by_row:
            raise Refused(f"row {r} isn't in this import")
        if not _joins(by_row[r]):
            raise Refused(f"row {r} ({by_row[r]['name']}) came in as its own technology, so there's no match to reopen")
    at, chosen = now_utc(), set(rows)
    decs = _import_decisions(imp)
    made_by = _made_by(imp, decs)
    # decisions recorded before they named their row now name it, so a later reopen needn't trace them again
    for d in decs:
        if d["_id"] in made_by and not (d.get("extra") or {}).get("row"):
            coll.update_one({"_id": d["_id"]}, {"$set": {"extra": {"import_id": import_id, "row": made_by[d["_id"]]}}})
    mine = {d["_id"] for d in decs}
    withdrawn = restored = 0
    for d in decs:
        if made_by.get(d["_id"]) not in chosen:
            continue
        coll.update_one({"_id": d["_id"]}, {"$set": {"active": False, "reopened_at": at, "reopened_tz": tz}})
        if d["active"]:
            withdrawn += 1
            restored += _restore_own(d, mine)
    targets = {by_row[r]["target_tech_id"] for r in rows}
    for r in rows:
        it = by_row[r]
        if it["group"] == "proposed":
            techs.update_one({"_id": it["target_tech_id"]}, {"$pull": {"aliases": it["name"].lower()}})
            it["decision"] = None
    joining = _joining(imp)
    for tech_id in targets:
        rec = (techs.find_one({"_id": tech_id}, {"imported": 1}) or {}).get("imported") or {}
        if rec.get("import_id") == import_id:
            its = joining.get(tech_id)
            techs.update_one({"_id": tech_id}, {"$set": {"imported": _record(imp, its, rec["at"], rec.get("tz"))}}
                             if its else {"$unset": {"imported": ""}})
    imp["status"] = "ready"
    imp["counts"] = counts(imp)
    imp.setdefault("reopened", []).append({"at": at, "tz": tz, "rows": sorted(rows), "withdrawn": withdrawn,
                                           "restored": restored})
    save(imp, set(rows))
    jobs.later(restamp, [by_row[r]["name"].lower() for r in rows])  # the pages that used the names taken back
    return {"reopened": len(rows), "values_withdrawn": withdrawn, "restored": restored}


# ── one technology by hand: a one-row import ─────────────────────────

def add_preview(name: str) -> dict:
    """What adding this name would do, before it's done: join a technology at
    no cost, offer one to confirm at no cost, or check it with the organize
    skill, with that call's cost."""
    group, hit, why = free_match(name)
    return {"outcome": group, "why": why,
            "target": {"tech_id": hit["_id"], **target_info([hit["_id"]]).get(hit["_id"], {})} if hit else None,
            "cost": match_cost([name.lower()]) if group == "unmatched" else None}


def template() -> bytes:
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "New Emerging Technologies"
    ws.append(HEADERS)
    ws.append(["Agentic AI", "AI", "AI Agents", "AI systems that plan and act on goals with little supervision.", "6",
               "Early Commercial"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
