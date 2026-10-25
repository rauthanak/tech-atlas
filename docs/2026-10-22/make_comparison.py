"""Builds the spreadsheet for Tata: the radar against Tata's catalog, before
and after the sources were widened on 22 Oct (13 → 86 sources). Reads the two
JSON files compare_catalog.py wrote (baseline.json, after.json).

    $SP/docvenv/bin/python make_comparison.py
"""

import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
OUT = HERE.parents[2] / "review-comments/10-22-2026/Tech Radar - radar vs Tata's catalog.xlsx"
NAVY = PatternFill("solid", fgColor="041A4A")
WHITE = Font(bold=True, color="FFFFFF")
ENOUGH = 3

before, after = (json.load(open(HERE / f)) for f in ("baseline.json", "after.json"))
B = {r["id"]: r for r in before["rows"]}
A = after["rows"]
cat_b = [r for r in before["rows"] if r["in_catalog"]]
cat_a = [r for r in A if r["in_catalog"]]


def bands(rows):
    return (sum(r["pages"] == 0 for r in rows), sum(0 < r["pages"] < ENOUGH for r in rows), sum(r["pages"] >= ENOUGH for r in rows))


def sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = NAVY, WHITE
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def why(r):
    if r["pages"] == 0:
        return f"None of the {len(after['sources_on'])} sources switched on has named it yet"
    if r["pages"] < ENOUGH:
        return f"Named on only {r['pages']} page{'s' if r['pages'] > 1 else ''}: too few to assess"
    return ""


wb = Workbook()
ws = wb.active
ws.title = "Summary"
lines = [
    ["Tech Radar against Tata's catalog, 22 October 2026"],
    ["What the radar finds that the catalog doesn't track, and how much evidence it has on what the catalog does. "
     "A page is one article, post, filing, paper or job posting that names the technology; "
     f"{ENOUGH} or more pages is enough to assess."],
    [],
    ["", "Before (13 sources)", f"After ({len(after['sources_on'])} sources)"],
    ["Technologies the radar knows", len(before["rows"]), len(A)],
    [f"Found by the radar, not in the catalog, with {ENOUGH}+ pages",
     sum(not r["in_catalog"] and r["pages"] >= ENOUGH for r in before["rows"]),
     sum(not r["in_catalog"] and r["pages"] >= ENOUGH for r in A)],
    [f"Catalog technologies ({len(cat_a)})", "", ""],
    ["   with no pages", *[bands(x)[0] for x in (cat_b, cat_a)]],
    [f"   with 1–{ENOUGH - 1} pages", *[bands(x)[1] for x in (cat_b, cat_a)]],
    [f"   with {ENOUGH}+ pages: enough to assess", *[bands(x)[2] for x in (cat_b, cat_a)]],
    [],
    ["Catalog technologies with no pages, by domain", "Before", "After", "In the catalog"],
]
for ln in lines:
    ws.append(ln)
tot = Counter(r["domain"] for r in cat_a)
zb = Counter(r["domain"] for r in cat_b if r["pages"] == 0)
za = Counter(r["domain"] for r in cat_a if r["pages"] == 0)
for d, n in tot.most_common():
    ws.append([d or "No domain", zb[d], za[d], n])
ws.append([])
ws.append(["Sources switched on after:", ", ".join(after["sources_on"])])
ws["A1"].font = Font(bold=True, size=14, color="041A4A")
for row in (4, 12):
    for c in ws[row]:
        c.font = Font(bold=True)
ws.column_dimensions["A"].width = 58
for col in "BCD":
    ws.column_dimensions[col].width = 20
ws["A2"].alignment = Alignment(wrap_text=True)
ws.row_dimensions[2].height = 45

found = sorted((r for r in A if not r["in_catalog"] and r["pages"] >= ENOUGH), key=lambda r: -r["pages"])
sheet(wb, "Not in your catalog", ["Technology", "Domain", "Commercial availability", "TRL", "Pages", "Pages before",
                                  "Sources", "Where it was found", "Radar status"],
      [[r["technology"], r["domain"], r["availability"], r["trl"], r["pages"], (B.get(r["id"]) or {}).get("pages", 0),
        len(r["sources"]), ", ".join(r["sources"]), r["status_reason"]] for r in found],
      [40, 22, 20, 6, 8, 12, 9, 60, 40])

cat_rows = sorted(cat_a, key=lambda r: (-r["pages"], r["technology"]))
sheet(wb, "Your catalog", ["Technology", "Domain", "Pages", "Pages before", "Sources", "Where it was found",
                           "Radar status", "Why little or nothing"],
      [[r["catalog_name"] or r["technology"], r["domain"], r["pages"], (B.get(r["id"]) or {}).get("pages", 0),
        len(r["sources"]), ", ".join(r["sources"]), r["status_reason"], why(r)] for r in cat_rows],
      [40, 22, 8, 12, 9, 60, 40, 50])
wb.save(OUT)
print("wrote", OUT, "| found:", len(found), "| catalog with pages now:", sum(r["pages"] > 0 for r in cat_a))
