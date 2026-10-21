"""Plan 7, step 7: version 2 of the signals and sources review.

Same rows as version 1 (5 Oct): the radar's signal types (Old), the 14 ideas the team sent (New) and our 3
suggestions. Version 2 marks plainly which of the team's ideas were new and added, which the radar already
had, and which aren't added yet; adds the column "Added to the code"; describes every source on its own tab;
and records the 5 Oct discussion. Names and on/off states come from the live radar (the script stops if a
row names something the radar doesn't have). Run from the Mac, with any Python that has openpyxl:
  python make_signals_xlsx.py "Signals and sources review - v2.xlsx"
"""

import json
import math
import sys
import urllib.request
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

API = "http://localhost:8100/api"
OUT = sys.argv[1]


def get(path):
    with urllib.request.urlopen(API + path, timeout=30) as r:
        return json.load(r)


SIG = {s["_id"]: s for s in get("/signals")["signals"]}
SRC = {s["_id"]: s for s in get("/sources")["sources"]}
NEW_TYPE = "approval-licence-or-certification"
RENAMED = {"spin-out": "Spin-out", "market-figure": "Market figure"}
GROUPS = {"@university_labs": ("university_labs", "university labs (MIT CSAIL, Stanford AI Lab, IISc, IIT Madras and others)"),
          "@conferences": ("conferences", "research conference sites (NeurIPS, SIGCOMM, USENIX NSDI and others)")}
MAILBOX = {"espacenet": "needs a free EPO account, which waits on the project mailbox",
           "sec_edgar_filings": "needs a contact email on every request, which waits on the project mailbox",
           "usajobs": "needs a free key registered with an email, which waits on the project mailbox"}
REFUSES = {"cisa", "fierce_network"}
TRIMMED = {"techstars", "3gpp", "etsi", "mplify", "tm_forum", "smpte", "rbi", "sebi", "amazon_science", "cisco_blogs", "tcl_news"}


def short_note(s):
    """A few words on why an off source is off, shown next to its name."""
    if s["_id"] in MAILBOX:
        return "waits on the project mailbox"
    if s["_id"] in REFUSES:
        return "refuses our collector"
    if s["cost"] == "licensed":
        return "needs a licence"
    if s["access"] == "none":
        return "no way in yet"
    return ""


def src_label(sid):
    if sid in GROUPS:
        cluster, text = GROUPS[sid]
        members = [s for s in SRC.values() if s["cluster"] == cluster]
        return f"{len(members)} {text}", any(s.get("enabled") for s in members)
    s = SRC[sid]  # KeyError: not in the radar's source list
    note = "" if s.get("enabled") else short_note(s)
    return s["name"] + (f" ({note})" if note else ""), bool(s.get("enabled"))


def sources(candidates=(), add=None, all_sources=False):
    """A row's sources: ids from the radar's list (or a group), plus sources worth adding."""
    return {"ids": list(candidates), "add": add, "all": all_sources}


def render(spec):
    if spec["all"]:
        return "Every source: each page read is checked for these"
    named = [src_label(i) for i in spec["ids"]]
    lines = []
    on = [n for n, enabled in named if enabled]
    off = [n for n, enabled in named if not enabled]
    if on:
        lines.append("Reading now: " + ", ".join(on))
    if off:
        lines.append("In our list, switched off: " + ", ".join(off))
    if spec["add"]:
        lines.append("Not in our list yet, worth adding: " + spec["add"])
    return "\n".join(lines)


OLD = "Old"
NEW_HAD, NEW_ADDED, NEW_NOT = "New – we already had it", "New – added in v2", "New – not added yet"
SUG_ADDED, SUG_NOT = "Suggested – added in v2", "Suggested – not added yet"
HAD_NO_CHANGE = "Already had it. No change."

FUND = ["techcrunch_venture", "hacker_news", "techstars", "inc42", "entrackr", "ycombinator", "crunchbase", "pitchbook", "dealroom", "tracxn"]
PILOT = ["light_reading", "gsma", "iot_analytics", "tm_forum", "verizon_innovation", "att_innovation", "rcr_wireless"]
REG = ["cert_in", "telegeography", "light_reading", "gsma", "rbi", "sebi", "meity", "bis_india", "enisa", "cisa"]
REG_ADD = "TRAI and DoT (India's telecom regulator and ministry)"
GRANT = ["openalex_grants", "darpa", "ukri", "dfg", "jst", "dst_india", "meity", "nsm_india", "nsf_news", "nsf_awards", "cordis"]
PATENT = ["espacenet", "uspto_ppubs", "uspto_odp", "wipo_patentscope"]
ANALYST = ["iot_analytics", "telegeography", "cx_today", "idc", "forrester", "cbinsights_research", "gartner", "mckinsey"]
COMPETITOR_ADD = ("competitors' newsrooms (Orange Business, BT, Lumen, NTT, Colt, Zscaler, Twilio, Sinch, Infobip, KORE) "
                  "and hyperscalers' newsrooms (AWS, Microsoft, Google Cloud)")

# (signal: a live id for Old rows, the team's wording for New rows, ours for Suggested; Old or New; Sources;
#  Already Configured, as the 5 Oct review found it, before any change; Added to the code)
ROWS = [
    ("funding-round", OLD, sources(FUND), "Yes, already in the radar.",
     "Changed in v2. It now says any stage, from pre-seed and seed to late stage, and counts a deep-tech company's first round as significant."),
    ("Pre-Seed / Seed funding round closed", NEW_HAD, sources(FUND),
     "Yes. Funding round already covered every round and noted its stage.",
     "We already had it. In v2, Funding round names pre-seed and seed outright, so small early rounds aren't treated as minor. "
     "Tested on a real page: Arq's $1.4M pre-seed round (The Quantum Insider)."),
    ("new-fund-or-theme", OLD, sources(["techcrunch_venture", "a16z", "sequoia", "lightspeed", "greylock", "index_ventures", "nfx"]), "Yes, already in the radar.", HAD_NO_CHANGE),
    ("Deep-tech startup incorporated or emerged from stealth", NEW_ADDED,
     sources(["techcrunch_venture", "hacker_news", "techstars", "ycombinator", "entrepreneur_first", "alchemist", "inc42", "entrackr", "crunchbase"]),
     "Partly. Spin-out covered companies spun out of labs, and Funding round caught stealth exits that announced money. "
     "A startup launching without money wasn't covered.",
     "Added in v2. Spin-out became \"Spin-out or startup launch\", which also covers a new deep-tech startup that is founded, incorporated or comes "
     "out of stealth. Tested on a real page: Eridu coming out of stealth (TechCrunch)."),
    ("acquisition", OLD, sources(["light_reading", "telegeography", "techcrunch_venture", "tcl_news", "fierce_network", "sec_edgar_filings"]), "Yes, already in the radar.", HAD_NO_CHANGE),
    ("Hyperscaler / Telecom competitors acquire, partner with a new firm/startup", NEW_HAD,
     sources(["light_reading", "telegeography", "techcrunch_venture", "cx_today", "tcl_news", "cisco_blogs", "fierce_network", "sec_edgar_filings"], add=COMPETITOR_ADD),
     "Yes. Acquisition or merger and Partnership catch the deal, and Competitor move (which names 20 TCL competitors) flags who did it.",
     "We already had it. In v2, Competitor move also includes hyperscalers (AWS, Microsoft, Google) when they move into TCL's markets."),
    ("ipo", OLD, sources(["techcrunch_venture", "sebi", "inc42", "entrackr", "sec_edgar_filings"]), "Yes, already in the radar.", HAD_NO_CHANGE),
    ("shutdown-or-failed-raise", OLD, sources(["techcrunch_venture", "hacker_news", "inc42", "entrackr"]), "Yes, already in the radar.", HAD_NO_CHANGE),
    ("Academic papers", NEW_HAD,
     sources(["arxiv", "openalex_works", "ieee_xplore", "acm_dl", "nature", "science", "ssrn", "@university_labs"]),
     "This is a kind of source, not an event. Papers were already read from arXiv; a paper is flagged when it reports a breakthrough or wins at a "
     "top conference.",
     "We already had it: arXiv is read on every collection. No change needed."),
    ("research-breakthrough", OLD,
     sources(["arxiv", "openalex_works", "amazon_science", "google_research", "microsoft_research", "meta_ai", "ibm_research", "nvidia_blog", "nature", "science"]),
     "Yes, already in the radar.", HAD_NO_CHANGE),
    ("top-conference-recognition", OLD, sources(["arxiv", "@conferences"]), "Yes, already in the radar.",
     "Already had it. No change. arXiv papers that name the conference they were accepted at are the main source."),
    ("Patent filings", NEW_NOT, sources(PATENT),
     "This is a kind of source, not an event. The signal \"Patent filing or grant\" existed, but no patent source was switched on, so patents were "
     "only caught when the news reported them.",
     "Not added yet. Reading patents needs a free account with the European Patent Office, which needs a project mailbox (see the Discussion tab). "
     "Until then, patents are caught only when the news reports them."),
    ("patent-filing-or-grant", OLD, sources(PATENT), "Yes, but no patent source was switched on.",
     "Already had it. No change, but still no patent source until the EPO account exists (see Discussion)."),
    ("Hyperscaler / Telecom competitors file a patent or research paper in a new domain", NEW_ADDED,
     sources(["arxiv", "openalex_works", "espacenet", "amazon_science", "google_research", "microsoft_research", "meta_ai", "nvidia_blog"],
             add="patent searches by applicant name (competitors and hyperscalers) on the EPO patent service"),
     "Partly. Patent filing or grant already highlighted patents by competitors and big technology companies, but a paper only counted if it was a "
     "breakthrough, and a company's first move into a new area wasn't defined.",
     "Added in v2. Competitor move now counts a company's first move into an area new to it (a first patent, paper, product, acquisition or team "
     "there), and hyperscalers moving into TCL's markets. Patents themselves still wait on the EPO account. In testing, the competitor label on a "
     "hyperscaler's move came up in some runs, not all; the move itself is always caught."),
    ("prototype-or-field-trial", OLD, sources(PILOT), "Yes, already in the radar.", HAD_NO_CHANGE),
    ("Pilot projects by Telco or large peer enterprises announced", NEW_HAD,
     sources(PILOT, add="operators' newsrooms (BT, Vodafone, Singtel, Airtel, Jio, NTT)"),
     "Yes, as Prototype or field trial: a demo, pilot or field trial with the parties named.",
     "We already had it, as Prototype or field trial. No change needed."),
    ("spin-out", OLD, sources(["techcrunch_venture", "hacker_news", "@university_labs", "ycombinator", "entrepreneur_first"]), "Yes, already in the radar.",
     "Changed in v2. Renamed from Spin-out, and it now also covers a new deep-tech startup that is founded, incorporated or comes out of stealth."),
    ("standard-milestone", OLD, sources(["ietf", "gsma", "3gpp", "etsi", "tm_forum", "mplify", "smpte", "itu", "iso", "nist_pubs", "irtf"]),
     "Yes, already in the radar.", HAD_NO_CHANGE),
    ("Standards body working group formed (ex: CAMARA project for 'Network APIs' managed by Linux Foundation)", NEW_HAD,
     sources(["ietf", "gsma", "linux_foundation", "3gpp", "etsi", "tm_forum"], add="the CAMARA project's and GSMA Open Gateway's own pages"),
     "Yes, as Standard milestone, which includes a standards body starting a work item or working group. CAMARA also counts as an Open-source milestone.",
     "We already had it, as Standard milestone. No change needed; IETF and GSMA are read on every collection."),
    ("open-source-milestone", OLD, sources(["hacker_news", "linux_foundation", "cncf", "apache", "lf_edge", "github_trending"]), "Yes, already in the radar.", HAD_NO_CHANGE),
    ("alliance-formed", OLD, sources(["gsma", "light_reading", "tm_forum", "linux_foundation", "wballiance"]), "Yes, already in the radar.", HAD_NO_CHANGE),
    ("partnership", OLD, sources(["light_reading", "gsma", "cx_today", "techcrunch_venture", "tcl_news", "cisco_blogs"]), "Yes, already in the radar.",
     "Changed in v2. It now also covers universities and governments, memorandums of understanding (MoUs) and joint research programmes."),
    ("MoU or research partnership announced", NEW_ADDED,
     sources(["light_reading", "gsma", "cx_today", "tcl_news", "cisco_blogs", "amazon_science"],
             add="press-release wires (PR Newswire, Business Wire) and EurekAlert! for university news"),
     "Partly. Partnership covered companies teaming up to build, integrate or sell; MoUs and research tie-ups with universities or governments "
     "weren't named.",
     "Added in v2. Partnership now covers universities and governments, MoUs and joint research. Tested on a real page: the 6G Flagship "
     "(University of Oulu) and Bharat 6G Alliance MoU."),
    ("product-launch", OLD,
     sources(["light_reading", "cx_today", "sports_video_group", "iot_analytics", "hacker_news", "cisco_blogs", "aws_reinvent", "google_cloud_next", "microsoft_ignite", "fierce_network"]),
     "Yes, already in the radar.", HAD_NO_CHANGE),
    ("competitor-move", OLD, sources(["light_reading", "telegeography", "cx_today", "sec_edgar_filings", "verizon_innovation", "att_innovation", "akamai_blog"],
                                     add=COMPETITOR_ADD),
     "Yes, already in the radar. It names 20 TCL competitors across the six fabrics.",
     "Changed in v2. It now also counts a company's first move into an area new to it, and hyperscalers (AWS, Microsoft, Google) moving into TCL's "
     "markets: enterprise connectivity (satellite included), cloud, security, customer communications, IoT or media. It's reported even when the "
     "same news is also a launch or a deal."),
    ("Competitor product launch or feature announcement", NEW_HAD,
     sources(["light_reading", "telegeography", "cx_today", "sec_edgar_filings", "akamai_blog", "fierce_network"], add=COMPETITOR_ADD),
     "Yes, as Product launch and Competitor move.",
     "We already had it, as Product launch and Competitor move. No change needed."),
    ("named-deployment", OLD, sources(["light_reading", "cx_today", "iot_analytics", "sports_video_group", "tcl_news", "cisco_blogs"]), "Yes, already in the radar.", HAD_NO_CHANGE),
    ("price-change", OLD, sources(["telegeography", "light_reading", "aws_reinvent", "google_cloud_next", "akamai_blog"]), "Yes, already in the radar.", HAD_NO_CHANGE),
    ("market-figure", OLD, sources(ANALYST), "Yes, already in the radar.",
     "Changed in v2. Renamed from Market figure, and it now also covers an analyst firm covering a technology for the first time, such as its "
     "first report on it or its first hype cycle, wave or quadrant."),
    ("Analyst firms publish first coverage", NEW_ADDED,
     sources(ANALYST, add="telecom analysts (Omdia, Analysys Mason, GSMA Intelligence, Dell'Oro)"),
     "Partly. Market figure caught analysts' numbers, but not a technology's first appearance in an analyst report, hype cycle or quadrant.",
     "Added in v2. Market figure became \"Market figure or analyst coverage\", which covers a technology's first appearance in an analyst report. "
     "Tested on a real page: Gartner's first-ever Magic Quadrant for cyberthreat intelligence (Group-IB, PR Newswire)."),
    ("regulation-or-mandate", OLD, sources(REG, add=REG_ADD), "Yes, already in the radar.", HAD_NO_CHANGE),
    ("Government or regulatory mandate", NEW_HAD, sources(REG, add=REG_ADD),
     "Yes, as Regulation or mandate: a law, rule, required standard, deadline or ban that forces or blocks adoption.",
     "We already had it, as Regulation or mandate. No change needed."),
    ("Regulatory approval or certification granted (ex: NIST certification for PQC algorithms)", NEW_ADDED,
     sources(["telegeography", "light_reading", "nist_pubs"], add="NIST's validation lists (CMVP and CAVP), TEC India (telecom equipment certification), FCC and Ofcom"),
     "Partly. Standard milestone caught NIST publishing a standard, but not a product or algorithm being certified or approved.",
     "Added in v2 as a new signal type, \"Approval, licence or certification\": a regulator or certification body approves, licenses or certifies "
     "a technology, product or provider. Tested on a real page: NIST's validation of Phio TX (Quantum XChange)."),
    ("government-programme-or-grant", OLD, sources(GRANT), "Yes, already in the radar.",
     "Already had it. No change. OpenAlex grants is switched on in v2, so new research grants are now read directly."),
    ("Funded university research program", NEW_HAD, sources(GRANT + ["@university_labs"]),
     "Yes for public money, as Government programme or grant. Labs funded by companies were only partly covered, through Hiring surge or new team.",
     "We already had it, as Government programme or grant. In v2, OpenAlex grants is switched on to catch new research grants directly."),
    ("public-tender", OLD, sources(["telegeography", "light_reading", "meity", "dst_india"], add="procurement portals: India's GeM and CPPP, and the EU's TED"),
     "Yes, already in the radar.", "Already had it. No change. Few sources carry tenders; the procurement portals would add more."),
    ("security-advisory", OLD, sources(["cert_in", "cisa", "enisa"]), "Yes, already in the radar.", HAD_NO_CHANGE),
    ("infrastructure-build", OLD, sources(["telegeography", "light_reading", "tcl_news"]), "Yes, already in the radar.", HAD_NO_CHANGE),
    ("major-outage", OLD, sources(["telegeography", "light_reading", "cloudflare_radar"]), "Yes, already in the radar.", HAD_NO_CHANGE),
    ("hiring-surge-or-new-team", OLD, sources(["job_boards", "usajobs"]), "Yes, already in the radar.",
     "Already had it. No change. Company job boards is switched on in v2."),
    ("employment-trend", OLD, sources(["usajobs", "stackoverflow_trends"]), "Yes, already in the radar.",
     "Already had it. No change, but no official jobs data until the USAJOBS key exists (see Discussion)."),
    ("developer-adoption-jump", OLD, sources(["hacker_news", "github_trending", "pypi", "npm_trends", "kaggle", "stackoverflow_trends"]), "Yes, already in the radar.", HAD_NO_CHANGE),
] + [(i, OLD, sources(all_sources=True), "Yes, already in the radar, as noise: pages like this are set aside.", HAD_NO_CHANGE)
     for i in ("single-job-ad", "event-promotion", "marketing-echo", "opinion-without-news", "routine-update", "award-marketing", "non-tech-corporate-news")] + [
    ("Customer demand: enterprises asking for, or tendering for, a technology", SUG_NOT,
     sources(add="TCL's own record of customer requests (internal); government tenders are already a signal type"),
     "No. TCL's framework starts a P3 prototype on customer demand, but the radar can't see what TCL's customers ask for.",
     "Not added yet. First a question for TCL: can it share a simple monthly summary of customer requests by technology? See the Discussion tab."),
    ("Spectrum or licence decision (e.g. satellite or private-5G spectrum)", SUG_ADDED,
     sources(["telegeography", "light_reading"], add="TRAI, DoT, FCC and Ofcom"),
     "Partly, through Regulation or mandate, which only covers rules that force or block adoption.",
     "Added in v2, inside the new signal type \"Approval, licence or certification\"."),
    ("Research or patent surge: a sharp rise in papers or patents on a topic", SUG_NOT, sources(["openalex_works", "espacenet"]),
     "No. A single page can't show a rise; it needs counting across collections.",
     "Not added yet. It needs a small build that counts papers and patents across collections. See the Discussion tab."),
]

# every live signal type but the new one is an Old row; the new type is reached through its two rows
old = [r[0] for r in ROWS if r[1] == OLD]
assert sorted(old) == sorted(set(SIG) - {NEW_TYPE}), set(SIG) ^ set(old)
assert NEW_TYPE in SIG and sum(r[1].startswith("New") for r in ROWS) == 14

ARIAL = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="041A4A")
MARK_FILL = {OLD: None, NEW_HAD: "DCEBFA", NEW_ADDED: "C9EBD5", NEW_NOT: "FBE3E3", SUG_ADDED: "C9EBD5", SUG_NOT: "FDF1D2"}
CONF_FILL = {"Yes": "E3F2EA", "Partly": "FDF1D2", "No": "FBE3E3", "This": "E8ECF1"}
ADDED_FILL = {"Added": "C9EBD5", "Changed": "C9EBD5", "We": "DCEBFA", "Already": "EEF1F4", "Not": "FBE3E3"}
WRAP = Alignment(wrap_text=True, vertical="top")
wb = Workbook()


def lines_needed(text, width):
    per = max(1, int(width * 1.2))  # Arial 10 fits a little over one character per unit of column width
    return sum(max(1, math.ceil(len(part) / per)) for part in str(text).split("\n"))


def table(ws, headers, widths, rows, fills=None):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 22
    for values in rows:
        ws.append(list(values))
        r = ws.max_row
        for c in ws[r]:
            c.font = Font(name=ARIAL, size=10)
            c.alignment = WRAP
        for col, color in (fills(values) if fills else {}).items():
            if color:
                ws.cell(r, col).fill = PatternFill("solid", fgColor=color)
        ws.row_dimensions[r].height = 13 * max(lines_needed(v, widths[i]) for i, v in enumerate(values)) + 5
    for i, w in enumerate(widths):
        ws.column_dimensions["ABCDEFGH"[i]].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{'ABCDEFGH'[len(headers) - 1]}{ws.max_row}"


def text_sheet(title, paragraphs, width=125):
    sh = wb.create_sheet(title)
    sh.column_dimensions["A"].width = width
    for text, size, bold in paragraphs:
        sh.append([text])
        c = sh.cell(sh.max_row, 1)
        c.font = Font(name=ARIAL, size=size, bold=bold)
        c.alignment = WRAP
        if text and size <= 10:
            sh.row_dimensions[sh.max_row].height = 13 * lines_needed(text, width) + 4
    return sh


# ---- Signals ----
ws = wb.active
ws.title = "Signals"
signal_rows = []
for sig, mark, src, conf, added in ROWS:
    name = (SIG[sig]["name"] + (f" (renamed from {RENAMED[sig]})" if sig in RENAMED else "")) if mark == OLD else sig
    signal_rows.append((name, mark, render(src), conf, added))
table(ws, ["Signal", "Old or New", "Sources", "Already Configured", "Added to the code"], [44, 18, 62, 52, 62], signal_rows,
      fills=lambda v: {2: MARK_FILL[v[1]], 4: CONF_FILL[v[3].split(" ")[0].rstrip(".,;:")], 5: ADDED_FILL[v[4].split(" ")[0].rstrip(".,;:")]})
for r in range(2, ws.max_row + 1):  # the team's ideas, and every changed row, in bold
    if ws.cell(r, 2).value != OLD or ws.cell(r, 5).value.startswith("Changed"):
        ws.cell(r, 1).font = Font(name=ARIAL, size=10, bold=True)

# ---- Sources: every source the Signals tab names, described in plain words ----
H_MEANING = {"H1": "H1 · in the market now", "H2": "H2 · emerging: money and talent moving", "H3": "H3 · early: research and standards"}
ON_DETAIL = {  # source: (what it gives us, signal types it mostly feeds, comment)
    "light_reading": ("Telecom industry news: operators' and vendors' launches, deals, deployments, pilots, prices and regulation.",
                      "Competitor move, Product launch, Named deployment, Prototype or field trial, Partnership, Acquisition or merger, Price change, "
                      "Regulation or mandate, Approval, licence or certification",
                      "The widest single source for TCL's Network, Cloud and IoT markets."),
    "telegeography": ("Subsea and land cables, data centres, outages and cable cuts, spectrum and licence awards, bandwidth prices and market figures.",
                      "Infrastructure build, Major outage or cable cut, Approval, licence or certification, Price change, Market figure or analyst coverage, "
                      "Regulation or mandate",
                      "Network infrastructure is TCL's core business."),
    "cx_today": ("Customer-communications news: contact centres, messaging and voice APIs, AI agents; competitors such as Twilio, Sinch and Infobip; "
                 "analyst rankings such as Magic Quadrants.",
                 "Competitor move, Product launch, Acquisition or merger, Partnership, Market figure or analyst coverage, Named deployment",
                 "Covers TCL's Interaction fabric (Kaleyra)."),
    "cert_in": ("India's national computer emergency response team: security advisories and binding security directions.",
                "Security advisory, Regulation or mandate",
                "The only security agency we can read: CISA refuses our collector."),
    "sports_video_group": ("Sports and live-media technology: broadcast and streaming launches, production deals and deployments.",
                           "Product launch, Partnership, Named deployment",
                           "Added as the 13th source so TCL's Media fabric has a source."),
    "techcrunch_venture": ("Startup money: funding rounds at every stage, new funds, acquisitions, IPOs, shutdowns, and startups coming out of stealth.",
                           "Funding round, New fund or investment theme, Acquisition or merger, IPO or listing, Shutdown or failed raise, "
                           "Spin-out or startup launch",
                           "Switched on in v2."),
    "hacker_news": ("What developers are talking about: new tools, open-source releases, startup launches and shutdowns. The radar follows each "
                    "post's link to the original article.",
                    "Developer adoption jump, Open-source milestone, Spin-out or startup launch, Product launch, Shutdown or failed raise", ""),
    "iot_analytics": ("IoT market research: market figures, analyst reports, deployments and launches.",
                      "Market figure or analyst coverage, Named deployment, Product launch",
                      "Covers TCL's IoT fabric (MOVE)."),
    "job_boards": ("How many roles 13 technology companies are hiring for, and how that changes: Cloudflare, Databricks, Anthropic, Zscaler, "
                   "Astranis, Netskope, IonQ, PsiQuantum, Fastly, Rigetti, OpenAI, Cohere and SandboxAQ.",
                   "Hiring surge or new team",
                   "Switched on in v2. One summary per company per collection; a single job ad is ignored by design."),
    "arxiv": ("New research papers in AI, networking and security.", "Research breakthrough, Top-conference recognition",
              "Covers the team's \"Academic papers\"."),
    "ietf": ("The body that writes internet standards: new working groups, adopted drafts and published standards.", "Standard milestone",
             "The clearest source for \"a standards body forms a working group\"."),
    "gsma": ("The mobile industry's body: standards and network APIs (Open Gateway, CAMARA), alliances, operator pilots and spectrum policy.",
             "Standard milestone, Alliance formed, Partnership, Prototype or field trial, Regulation or mandate", ""),
    "openalex_grants": ("New research grants from a free global research database, on 8 saved topics: quantum key distribution, post-quantum "
                        "cryptography, coherent optical networks, network slicing, LEO satellite broadband, zero trust network access, edge AI "
                        "inference and network digital twins.",
                        "Government programme or grant",
                        "Switched on in v2. Covers the team's \"Funded university research program\". The topics can be changed in Setup."),
}
HOW = {"crawl": "Website (and its news feed, where it has one), up to 60 pages per collection",
       "query": "Free data service (API), with saved searches", "none": "No way in yet"}
CLUSTER = {
    "analysts_trade_press": "Industry analysts and trade press: market figures, analyst reports and industry news.",
    "company_filings": "Companies' official US filings: results, deals and launches, and Form D for startup fundraising.",
    "corporate_research_vendor_events": "Big-tech research labs, vendor blogs and conferences: research results and product launches.",
    "government_funding": "Governments, regulators and research funders: programmes, grants, regulation and security advisories.",
    "standards_foundations": "Standards bodies and open-source foundations: new standards, working groups and projects.",
    "talent": "Job postings: who is hiring for which skills.",
    "vc_startup_developer": "Startups, investors and developers: funding, new funds, launches and developer activity.",
    "conferences": "Top research conferences: accepted papers and awards.",
    "patents": "Patent offices: patents filed and granted.",
    "preprints_journals": "Research papers and journals.",
    "university_labs": "University research labs: new research and spin-outs.",
}


def off_comment(s):
    sid = s["_id"]
    if sid in MAILBOX:
        return MAILBOX[sid][0].upper() + MAILBOX[sid][1:] + " (see Discussion)."
    if sid in REFUSES:
        return "Refuses our collector (HTTP 403), so it can't be read."
    if sid == "meity":
        return "Its website didn't load from here when checked on 4 Oct."
    if s["cost"] == "licensed":
        return "Needs a paid licence; it would come from TCL."
    if s["access"] == "none":
        return "No way in yet: there's no feed or free service the radar can use."
    if sid in TRIMMED:
        return "Switched off in v2 to keep the run to 13 sources; can be switched back on in Setup."
    return "Switched off to keep the run to 13 sources; can be switched back on in Setup."


# which signal types name each source (the new type through the team's certification row)
mentioned = {i for _, _, spec, _, _ in ROWS for i in spec["ids"] if not i.startswith("@")}
feeds_of = {}
for sig, mark, spec, _, _ in ROWS:
    sig_id = sig if mark == OLD else (NEW_TYPE if sig.startswith("Regulatory approval") else None)
    for i in spec["ids"] if sig_id else []:
        if SIG[sig_id]["name"] not in feeds_of.setdefault(i, []):
            feeds_of[i].append(SIG[sig_id]["name"])

src_rows = []
order = {"H1": 1, "H2": 2, "H3": 3}
for s in sorted((SRC[i] for i in mentioned | set(ON_DETAIL)), key=lambda s: (not s.get("enabled"), order[s["horizon"]], s["name"].lower())):
    on = bool(s.get("enabled"))
    if on:
        what, sigs, comment = ON_DETAIL[s["_id"]]
    else:
        what, sigs, comment = CLUSTER[s["cluster"]], ", ".join(feeds_of.get(s["_id"], [])) or "—", off_comment(s)
    how = HOW[s["access"]] + ("" if s["cost"] == "free" else " · " + {"free_tier": "free tier", "licensed": "licensed"}[s["cost"]])
    src_rows.append((s["name"], H_MEANING[s["horizon"]], "Reading now" if on else "In our list, switched off", how, what, sigs, comment))
for token, (cluster, text) in GROUPS.items():  # the two groups the Signals tab names as one entry
    members = sorted((m for m in SRC.values() if m["cluster"] == cluster), key=lambda m: m["name"].lower())
    src_rows.append((f"{len(members)} {text.split(' (')[0]}: " + ", ".join(m["name"] for m in members), H_MEANING[members[0]["horizon"]],
                     "In our list, switched off", HOW["crawl"], CLUSTER[cluster], ", ".join(feeds_of.get(token, [])) or "—",
                     "Switched off to keep the run to 13 sources; can be switched back on in Setup."))
ADD_ROWS = [
    ("Competitors' newsrooms: Orange Business, BT, Lumen, NTT, Colt, Zscaler, Twilio, Sinch, Infobip, KORE", "H1", "Competitors' own announcements.",
     "Competitor move, Product launch, Partnership", "Would make the competitor signals much stronger."),
    ("Hyperscalers' newsrooms: AWS, Microsoft, Google Cloud", "H1", "The hyperscalers' own launches and moves into TCL's markets.",
     "Competitor move, Product launch", "Pairs with the v2 change that counts hyperscalers in Competitor move."),
    ("Operators' newsrooms: BT, Vodafone, Singtel, Airtel, Jio, NTT", "H1", "Operators' pilots, trials and deployments.",
     "Prototype or field trial, Named deployment", "Covers the team's \"Pilot projects by telcos\" more directly."),
    ("Telecom analysts: Omdia, Analysys Mason, GSMA Intelligence, Dell'Oro", "H1", "Telecom market figures and analysts' new reports.",
     "Market figure or analyst coverage", "Mostly paywalled; their free press pages carry the headlines."),
    ("TRAI and DoT (India's telecom regulator and ministry)", "H1", "Indian telecom rules, consultations, spectrum and licences.",
     "Regulation or mandate, Approval, licence or certification", "TCL's home regulator."),
    ("FCC and Ofcom (US and UK regulators)", "H1", "Spectrum decisions, licences and equipment approvals.",
     "Approval, licence or certification, Regulation or mandate", ""),
    ("NIST's validation lists (CMVP and CAVP)", "H3", "Which cryptographic products and algorithms NIST has validated, including post-quantum ones.",
     "Approval, licence or certification", "The team's example of a certification signal."),
    ("TEC India", "H3", "Indian telecom equipment certification.", "Approval, licence or certification", ""),
    ("CAMARA project and GSMA Open Gateway pages", "H3", "Network-API standards and the operators adopting them.",
     "Standard milestone, Open-source milestone", "The team's example of a standards working group."),
    ("Press-release wires: PR Newswire, Business Wire", "H1", "Companies' press releases, including MoUs and launches.",
     "Partnership, Product launch", "Very noisy at 60 pages a collection; best used with a filtered feed."),
    ("EurekAlert! (university research news)", "H3", "Universities' research news and partnerships.", "Partnership, Research breakthrough", ""),
    ("Procurement portals: India's GeM and CPPP, the EU's TED", "H1", "Public bodies buying technology.", "Public tender",
     "Few current sources carry tenders."),
]
for name, h, what, sigs, comment in ADD_ROWS:
    src_rows.append((name, H_MEANING[h], "Not in our list yet, worth adding", "—", what, sigs, comment))
sw = wb.create_sheet("Sources")
STATUS_FILL = {"Reading now": "C9EBD5", "In our list, switched off": "EEF1F4", "Not in our list yet, worth adding": "FDF1D2"}
table(sw, ["Source", "Horizon", "Status", "How the radar reads it", "What it gives us", "Signal types it mostly feeds", "Comment"],
      [34, 22, 18, 30, 50, 44, 44], src_rows, fills=lambda v: {3: STATUS_FILL[v[2]]})

# ---- Notes ----
on_by_h = {h: [s["name"] for s in SRC.values() if s.get("enabled") and s["horizon"] == h] for h in ("H1", "H2", "H3")}
team = [r for r in ROWS if r[1].startswith("New")]
count = {m: sum(r[1] == m for r in ROWS) for m in (NEW_ADDED, NEW_HAD, NEW_NOT, SUG_ADDED, SUG_NOT)}
H, T = 11, 10
text_sheet("Notes", [
    ("Signals and sources review, version 2", 14, True),
    (f"Made on {date.today().day} {date.today():%b %Y}, after the changes in Plan 7 went into the radar (tech-atlas/docs/2026-10-05/1). Version 1, from 5 Oct, "
     "is kept in replaced/.", T, False),
    ("", T, False),
    ("In short", H, True),
    (f"The team sent {len(team)} ideas. {count[NEW_ADDED]} were new to the radar and are now added; {count[NEW_HAD]} the radar already had; "
     f"{count[NEW_NOT]} isn't added yet (patents, which need an account with the European Patent Office). Of our own 3 suggestions, "
     f"{count[SUG_ADDED]} is added and {count[SUG_NOT]} are not yet.", T, False),
    (f"The radar now has {len(SIG)} signal types and reads {sum(map(len, on_by_h.values()))} sources, up to 60 pages each per collection.", T, False),
    ("Sources read now: H1 " + ", ".join(on_by_h["H1"]) + ". H2 " + ", ".join(on_by_h["H2"]) + ". H3 " + ", ".join(on_by_h["H3"]) + ".", T, False),
    ("Each change was tested on real pages before it went in. Everything collected before 5 Oct was cleared, and the radar is being rebuilt from "
     "a fresh collection on these 13 sources. Whether every signal type shows up depends on what those sources publish.", T, False),
    ("", T, False),
    ("How to read the Signals tab", H, True),
    ("Old or New says where each row came from:", T, True),
    ("  Old: a signal type the radar already had before this review. Renamed ones show both names.", T, False),
    ("  New – added in v2 (green): sent by the team, and new to the radar. It's now in the radar.", T, False),
    ("  New – we already had it (blue): sent by the team, but the radar already covered it.", T, False),
    ("  New – not added yet (red): sent by the team; the last column says why it isn't in yet.", T, False),
    ("  Suggested: our own ideas, marked added or not added yet the same way.", T, False),
    ("Sources: \"Reading now\" means switched on. \"In our list, switched off\" means the radar knows it but doesn't read it now, with the reason "
     "in brackets. \"Not in our list yet\" is a suggestion. The Sources tab describes every source in plain words.", T, False),
    ("Already Configured: how the radar stood on 5 Oct, before any change, so you can compare before and after.", T, False),
    ("Added to the code: what happened. \"Added in v2\" or \"Changed in v2\" (green) means it's now in the radar; \"We already had it\" (blue) or "
     "\"Already had it\" (grey) means no change was needed; \"Not added yet\" (red) says why.", T, False),
    ("Bold signal names are the team's ideas and the types that changed in v2.", T, False),
    ("", T, False),
    ("Good to know", H, True),
    ("A signal type isn't tied to particular sources: every page from every source that's switched on is checked against every signal type. "
     "The Sources column shows where each type usually turns up.", T, False),
    ("The new signal type, \"Approval, licence or certification\", has no row of its own: it answers the team's \"Regulatory approval or "
     "certification granted\" and our spectrum suggestion, so it's shown on those two rows.", T, False),
])

# ---- Discussion ----
text_sheet("Discussion", [
    ("Points and questions from the 5 Oct discussion", 14, True),
    ("", T, False),
    ("1. A project mailbox, for three data sources", H, True),
    ("What it is: one email address that belongs to the radar itself, not to a person. The team owns it now and hands it to TCL with the radar "
     "in December.", T, False),
    ("Why it's needed: three free data sources only work with an email address.", T, False),
    ("  EPO patents (the European Patent Office's free data service) needs a free account, which gives a key. It's our only working way to read "
     "patents (the US and world patent offices have no way in yet), so until then the team's \"Patent filings\" stays uncovered.", T, False),
    ("  SEC EDGAR (US company filings: Form D shows startups raising money; 8-K shows deals and launches) asks for a contact email on every request.", T, False),
    ("  USAJOBS (US government job postings, for employment trends) needs a free key registered with an email, which is also sent with every request.", T, False),
    ("Why not a personal email: a team member's own email never goes to outside services, and accounts tied to a student's address would leave "
     "with the students in December.", T, False),
    ("Two ways to get one: (a) a new free mailbox now (Gmail or Outlook), handed to TCL in December. Recommended, because it holds nothing up. "
     "(b) A shared address from TCL. Better for the handover, but slower.", T, False),
    ("Who does what: the team creates the mailbox and signs up at EPO and USAJOBS (Claude can't create accounts). The keys go into the radar's "
     "private settings file, never into the seeding data or git. Claude then switches the three sources on, tests them, and they join the next collection.", T, False),
    ("Status: open.", T, True),
    ("", T, False),
    ("2. Customer demand: a question for TCL", H, True),
    ("Why it matters: TCL's own framework sets P3 as \"prototype or MVP on customer demand\", so customers asking for a technology is a reason to "
     "act. The radar reads only public sources, so it can't see that demand. It sits inside TCL: the RFPs and presales requests TCL receives, "
     "and account managers' notes. Government tenders are already a signal type, but private companies' RFPs aren't public.", T, False),
    ("The question for TCL: Can TCL share a simple monthly summary of customer requests by technology (technology, fabric, region, count), with "
     "no customer names?", T, True),
    ("If TCL says yes: the radar could read that summary and flag customer demand next to the public signals. That's a small build, later. "
     "If TCL says no: nothing changes, and TCL keeps applying its P3 rule by judgement.", T, False),
    ("Status: to raise with TCL. The team called it a very valid question.", T, True),
    ("", T, False),
    ("3. Decisions made on 5 Oct", H, True),
    ("Signals: widen Partnership, Competitor move, Spin-out (now Spin-out or startup launch) and Market figure (now Market figure or analyst "
     "coverage); reword Funding round to say any stage; add one new type, Approval, licence or certification. Hyperscalers are counted in "
     "Competitor move: \"who are we to stop them if they think their competitors are hyperscalers\".", T, False),
    ("Sources: 13 switched on, 4 for each horizon plus Sports Video Group for Media, up to 60 pages each per collection. Run cost isn't a concern.", T, False),
    ("Data: everything collected before 5 Oct was cleared and the radar is rebuilt from a fresh collection. The backup is in "
     "tech-atlas/docs/2026-10-05/1/replaced/db-before-plan7.archive.gz and puts everything back with one command.", T, False),
    ("Model: gpt-6-luna from 5 Oct, cheaper than gpt-5.6-luna. On it, the radar's agents also work again; they had been failing at their first "
     "call since 4 Oct.", T, False),
    ("Assessments: technologies that appear on 3 or more pages get scored and a Point of View.", T, False),
    ("", T, False),
    ("4. Known gaps, and what would close them", H, True),
    ("Patents: no source until the EPO account exists (point 1).", T, False),
    ("Employment trends: no official jobs data until the USAJOBS key exists (point 1).", T, False),
    ("Paid sources (Crunchbase, PitchBook, Dealroom, Tracxn, Gartner): need TCL's licences.", T, False),
    ("A rise in papers or patents on a topic: needs a small build that counts them across collections, once the paper and patent sources have run "
     "a few times.", T, False),
    ("Hyperscalers as competitors: their moves are always caught (as a launch, an approval or a partnership), but the model adds the Competitor "
     "move label only some of the time. Worth checking after each collection.", T, False),
    ("Sources that refuse our collector: CISA and Fierce Network.", T, False),
    ("", T, False),
    ("5. Open questions", H, True),
    ("For the team: the project mailbox (point 1).", T, False),
    ("For the team: should a collection read the three horizons at the same time? Today they run one after another, about 20 minutes each. "
     "Side by side, a collection would take about 20 minutes, plus organizing the technologies. It's a small code change and one setting.", T, False),
    ("For TCL: the customer-demand summary (point 2). And the settings that still carry our proposal: the priority in each grid cell (TCL's chart "
     "shows P5 where ours shows P4), which priorities should trigger action, the horizon map, each fabric's TCL-today profile, the market-size "
     "anchors, a fuller description of TCL's portfolio, and whether relevance differs by business unit or region.", T, False),
])

wb.save(OUT)
print("wrote", OUT, "| rows:", len(ROWS), {m: sum(r[1] == m for r in ROWS) for m in MARK_FILL}, "| sources tab:", len(src_rows))
