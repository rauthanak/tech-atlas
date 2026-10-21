"""Needs python-docx (pip install python-docx in any venv). Builds the client change log: Tech Radar - changes after the midpoint review.docx.

For each change: where it came from (the email, document or page), what changed, why, before → after in numbers,
a workflow diagram where it helps (make_diagrams.py), and the screens before and after.
Re-run after each change: update ITEMS, N (the numbers, from the live radar) and HISTORY."""
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path("/Users/akki/Documents/Claude/Code/Tech-Radar/review-comments/10-20-2026")
BEFORE = ROOT / "screenshots/before"
AFTER_DESIGN = ROOT / "screenshots/after-design"
AFTER_BUILT = ROOT / "screenshots/after"
DIAGRAMS = ROOT / "screenshots/diagrams"
OUT = ROOT / "Tech Radar - changes after the midpoint review.docx"
NAVY = RGBColor(0x04, 0x1A, 0x4A)
MUTED = RGBColor(0x49, 0x4A, 0x52)

VERSION = "Version 4 · 21 October 2026"
HISTORY = [("1", "20 Oct 2026", "First version: every change agreed after the midpoint review, with today's screens and design mock-ups."),
           ("2", "20 Oct 2026", "Change 0 built: technologies say how solid their assessment is, a technology's result in each fabric stays "
                                "with the technology, plain names on every screen. Real screens replace its mock-up."),
           ("3", "21 Oct 2026", "Changes 1–13 built. Real screens replace their mock-ups. Change 5 tried both ways first; two measures moved to "
                                "per fabric as a result. Tata's list taken up to the review step."),
           ("4", "21 Oct 2026", "Rearranged so each change says where it came from, with before → after numbers and workflow diagrams. "
                                "Numbers now come from the full run of 21 Oct: every technology assessed again and every Point of View "
                                "rewritten with numbered sources. Tata's list imported (885 added, 115 joined). Frontier gains Tata's "
                                "stage buckets.")]

# The numbers in the text, from the live radar on 21 Oct (6 Oct is the run before these changes).
N = dict(
    old_techs=147, old_calls="4,474", old_scored=50, old_unscored=50,  # the two measures Tata enters left out
    techs_all="2,490", calls_full=772, techs_full=248, per_tech=3, scored_full=84, unscored_full=16, reused_pct=25,
    calls_recorded="3,623", classified="2,490", nascent=159, proposed_subdomains=246, skills=18,
    placed=188, held=61, ready=1, thin="1,357", noev=883, off=0,
    povs=250, pov_sources=250, avg_sources=9,
    fr_placed=240, fr_techs=114, fr_mature=85,
)

# Where the changes came from; each change cites these by their code.
SOURCES = [
    ("S1", "Hasan's email “Midpoint Review Summary (Oct 13)”",
     "The agreed record of the 13 Oct review, with the replies from Abhishek, Nick and Ankur (Ankur's three tests, 17 Oct)."),
    ("S2", "Email “Mid Point Summary – Zoom”", "Zoom's summary of the 13 Oct meeting, and Nick's note of 19 Oct."),
    ("S3", "The review deck (Tata_Tech_Radar_Midpoint_Review_Oct13.pdf)", "19 pages; the page is given each time."),
    ("S4", "Nick's note (Tata_Finance_Value_Levers_v2.docx)", "Value levers, blockers to test, and finance, revised 20 Oct."),
    ("S5", "Tata's list (Emerging-technologies October 2026.xlsx)",
     "Attached by Abhishek to S1: 1,000 rows in the 'New Emerging Technologies' tab."),
    ("S6", "Our working sessions on the comments, 20 Oct", "Where we went through S1–S5 point by point and agreed how to build each."),
]

# Built screens, in reading order, with their captions: item-N.png or item-N-<screen>.png
AFTER_CAPTIONS = {
    "item-0-technologies.png": "After: 'Tracked' replaces 'Up to date'. Each row shows its priority and how many of its measures had "
                               "enough evidence to score.",
    "item-0-technology.png": "After: a technology's page says how solid its assessment is, and each fabric shows its own count.",
    "item-0-skills.png": "After: skills by plain name, in the order the radar uses them; code names only under Technical details.",
    "item-0-scoring-tab.png": "After: every Scoring method tab has a title and a one-line purpose.",
    "item-0-sources.png": "After: Sources opens on the ones in use; 'Show all' lists the rest.",
    "item-1-bulk-bar.png": "After: tick rows, and a bar acts on all of them: switch off, switch on, move, or assess with the cost shown first.",
    "item-1-switched-off.png": "After: a switched-off technology says when, in the time zone it was switched off in; its pages are still collected.",
    "item-2-written-time.png": "After: a Point of View written from India shows India's time; older ones show US Central.",
    "item-3-skill-map.png": "After: every skill in the step that uses it, calls that carry several skills as dashed boxes, and what each used last time.",
    "item-4-operations.png": "After: each run in Operations' history says how many tokens it used, and how many were reused.",
    "item-4-collection.png": "After: a collection's page shows its AI use: calls, tokens sent, reused and written back.",
    "item-5-engine.png": "After: Weights & Grid chooses how an assessment calls the AI: stacked (recommended) or one measure per call.",
    "item-5-stacked.png": "After: a skill's page shows how it's sent inside a stacked call, its own parts highlighted, the reused parts first.",
    "item-6-fabric-tab.png": "After: one tab per fabric on the technology's page: the priority (override with a reason), and the measures that "
                             "need a person first: score by hand, accept the gap, Tata's entries, find evidence.",
    "item-6-briefing.png": "After: the Briefing opens on the two to-do lists, Before scoring and Before publishing, counted by technology.",
    "item-7-technologies.png": "After: views as Tata thinks of them, filters, and for each technology where it is with the reason, its "
                               "classification and its evidence.",
    "item-8-answers.png": "After: four answers on every technology, each saying where it came from, with Why? (reason and pages) and Change.",
    "item-8-classification.png": "After: Setup › Classification: Tata's domains and sub-domains, with the AI's proposed sub-domains to accept, "
                                 "merge or rename.",
    "item-9-nascent.png": "After: a possibly nascent technology, with the reason and the quote that flagged it; a person answers yes or no.",
    "item-10-find-evidence.png": "After: Find evidence and Add a link wherever evidence is thin, with the cost and today's search budget first.",
    "item-11-import.png": "After: Tata's list, imported: 88 matched exactly, 27 proposed matches accepted, 885 added. Nothing "
                          "was saved until it was confirmed.",
    "item-12-value-tab.png": "After: where the value lands in a fabric: use cases with Nick's lever chain, the path to scale, and the finance card.",
    "item-12-pov.png": "After: the Point of View leads with the verdict, then priority, maturity, availability and evidence; numbered "
                       "sources at the end.",
    "item-13-frontier.png": "After: Frontier at its own address: the 2×2 by fabric and by Tata's stage (emerging by default, Mature on "
                            "request), the table, and what isn't placed yet.",
}

# status: Planned · In progress · Built · Live
ITEMS = [
    dict(n="0", title="Clearer screens, first fixes", status="Built",
         source=[("S6", "a walk through the portal as a first-time user found 11 places where it confuses; the quick ones were "
                        "fixed first, the rest folded into the changes below")],
         what=["'Up to date' is now 'Tracked', and no longer suggests 'complete'. Every technology shows its priority and how many "
               "of its measures had enough evidence to score; amber marks those with no priority, or under half scored.",
               "A technology's result in one fabric opens on the technology's own page, as a tab, instead of under Operations, so "
               "you don't lose your place.",
               "Plain names replace internal codes: skills and measures by name, evidence by collection date; codes only under "
               "Technical details.",
               "Every Scoring method page has a title and a one-line purpose; the sources page opens on the sources in use."],
         why="Users said the portal is still confusing, and the changes below add more to it, so it had to be clear first.",
         metrics=[("Technologies shown as 'Up to date', in green", "146, although 58% of their measures were unscored",
                   "None: each says 'Tracked', with its priority and 'x of y measures scored'"),
                  ("Where a technology's result in a fabric opens", "Under Operations (you lose your place)",
                   "A tab on the technology's own page")],
         before=[("02-technologies.png", "Before: 146 technologies said 'Up to date', in green, although most measures were unscored."),
                 ("08-assessment-agentic-ai-cloud.png", "Before: a technology's result opened under Operations, with measure codes "
                  "and 'unscorable'.")]),
    dict(n="1", title="Switch a technology off from the list", status="Built",
         source=[("S1", "'Things to build': switch a technology off from the technology list, committed at the review"),
                 ("S6", "how it behaves: its pages are still collected, and it asks for a look when new pages arrive")],
         what=["An on/off switch on every row of the technology list; tick several rows to switch them off or on at once.",
               "A switched-off technology isn't assessed or shown to leaders, but its pages are still collected, so nothing is lost. "
               "It says when it was switched off.",
               "If it gathers new pages after being switched off (10 by default, a setting), the list asks for a look."],
         why="Agreed at the midpoint review.",
         metrics=[("Switching off 10 technologies", "Open each one, then Manage, then the switch: 10 pages",
                   "Tick 10 rows, one click")],
         before=[("02-technologies.png", "Before: the switch existed only inside each technology's page.")]),
    dict(n="2", title="When a Point of View was written: date, time and time zone", status="Built",
         source=[("S6", "asked when we went through the comments"),
                 ("S3", "page 6: what leaders see is dated ('published 6 Oct')")],
         what=["Each Point of View shows the date and time it was written, in the time zone of the person who wrote it."],
         why="A reader needs to know how fresh a Point of View is, and readers sit in India and the US.",
         metrics=[("A Point of View's written time", "'Oct 6'", "'Oct 21, 2026, 7:45 AM GMT+5:30'")],
         before=[("05-pov-agentic-ai.png", "Before: 'Written Oct 6', with no time.")]),
    dict(n="3", title="One picture of where each AI instruction is used", status="Built",
         source=[("S1", "action item for Akshay: a diagram of where each prompt is used")],
         what=["A page shows every instruction the AI follows (we call them skills) in the step that uses it, the calls that carry "
               "several skills together, and what each used in its last run.",
               "It prints on one page, so it can be shared."],
         why="So anyone can see what the AI is asked to do, where, and at what cost.",
         metrics=[("Skills shown together, in the order they run", "None: one page per skill", "{skills}, on one page")],
         before=[("09-skills.png", "Before: a list of skills, each explained on its own page.")]),
    dict(n="4", title="AI usage recorded for every call", status="Built",
         source=[("S6", "our plan to measure the cost of every change before and after, and keep it under control")],
         what=["Every AI call records the tokens it sent, reused and wrote back, with its step, run and technology. Operations, each "
               "run's page and each skill's page show them.",
               "The fixed opening of a call (its instructions and evidence) is marked for reuse, so the next call that starts the "
               "same way pays a tenth of the price for that part.",
               "Each skill has a switch for whether the AI may reason before it answers. Reasoning is paid for but never shown, so "
               "it's on only where it helps: scoring, value levers and the Point of View."],
         why="Without it, nobody could say what a run or a change costs.",
         metrics=[("AI calls with their tokens recorded", "None", "Every call: {calls_recorded} since 20 Oct"),
                  ("Tokens sent that were reused at a tenth of the price (21 Oct run)", "Nothing reused", "{reused_pct}%")],
         before=[("11-operations.png", "Before: run history showed the number of AI calls, not their size or cost.")]),
    dict(n="5", title="A leaner assessment", status="Built",
         source=[("S6", "why the assessment called the AI about 30 times per technology, and the plan to change it")],
         what=["Before, each technology got one AI call per measure per fabric, about 30 in all, each sending all of its evidence again.",
               "Now one call answers the questions about the technology as a whole (maturity, market size, growth, which fabrics it "
               "matters to, metrics to watch, and its domain, TRL and availability), and one call per relevant fabric answers that "
               "fabric's questions. The evidence is sent once and reused.",
               "Both ways were tried on the same 6 technologies, with the same AI and the same pages, before switching. Commercial "
               "adoption and use-case breadth came out too high when judged for the technology as a whole, so they're judged per fabric.",
               "Each assessment may now read up to 500,000 tokens of evidence (roughly a 750-page book), up from 300,000; "
               "the AI's limit is 1,000,000."],
         why="The same evidence was being sent about 30 times per technology, and the same maturity question asked once per fabric.",
         metrics=[("AI calls per technology", "About 30 ({old_calls} for {old_techs} on 6 Oct)",
                   "About {per_tech} ({calls_full} for {techs_full} on 21 Oct)"),
                  ("The same 6 technologies, both ways: AI calls", "204", "19"),
                  ("The same 6: tokens at the full price", "4,287,124", "367,485 (about 12 times fewer)"),
                  ("The same 6: measures the evidence could answer", "78%", "93%"),
                  ("Every assessed technology: measures the evidence could answer (the two Tata enters left out)",
                   "{old_scored}% (6 Oct, with the earlier AI model)", "{scored_full}% (21 Oct)")],
         diagram=("engine", "Before, every measure was its own call; now one call per technology and one per fabric."),
         before=[("08-assessment-agentic-ai-cloud.png", "Before: one assessment page per technology and fabric, built from one call per measure.")]),
    dict(n="6", title="The innovation team's decisions before publishing (Gate 2)", status="Built",
         source=[("S3", "pages 9, 11 and 14: Gate 2 overrides are audited and carry a reason; 'handle unscored measures'"),
                 ("S2", "Hasan: place a technology by hand when the scores can't place it"),
                 ("S1", "the not-yet-placed list stays visible beside the grid")],
         what=["Each fabric has a tab on the technology's page with its priority. A person can override it, with a reason; the "
               "computed priority stays beside it, and Undo restores it.",
               "Measures the evidence couldn't score come first, each with three ways to settle it: score it by hand (with a reason), "
               "accept the gap, or find evidence.",
               "Disruption threat and Serviceable capture are entered by Tata: public pages almost never answer them.",
               "The Briefing opens on two to-do lists: Before scoring (Gate 1) and Before publishing (Gate 2).",
               "A person's choices are kept, and a new run never overwrites them."],
         why="'An unowned score is overruled' (S3, page 3). On 6 Oct, 58% of measures couldn't be scored from public pages, "
             "and nothing could be done about them.",
         metrics=[("Measures the evidence couldn't score (the two Tata enters left out)",
                   "{old_unscored}% (6 Oct), with nothing to do about them",
                   "{unscored_full}% (21 Oct), each with score by hand, accept the gap, or find evidence"),
                  ("Disruption threat and Serviceable capture", "Asked of the AI: 98% unscored", "Entered by Tata"),
                  ("Priority set by a person", "Not possible", "Override with a reason; the computed one stays visible")],
         diagram=("process", "The whole process now: the two gates are where people decide."),
         before=[("08-assessment-agentic-ai-cloud.png", "Before: unscored measures said what was missing, but nothing could be done about them.")]),
    dict(n="7", title="Technologies organised the way Tata thinks about them", status="Built",
         source=[("S1", "Shashank: mature and emerging views, nascent technologies first; raise the evidence threshold"),
                 ("S6", "'Tracked' means assessed; organise the list by assessed, not assessed and why, emerging, mature but not on "
                        "Tata's radar")],
         what=["Every technology says where it is and why: placed on the grid, held beside it, ready to assess, thin evidence, "
               "no evidence yet, or switched off.",
               "Ready-made views: Tracked · Not assessed · Emerging · Emerging and new to Tata · Mature but not on Tata's radar · "
               "Possible nascent · Imported with little evidence · Switched off.",
               "Filters by commercial availability, domain, TRL, evidence and fabric; act on several rows at once."],
         why="Mature technologies shouldn't crowd out emerging ones, and a technology that isn't assessed should say why.",
         metrics=[("Technologies that say why they're where they are", "None: 1,003 just said 'Not assessed'",
                   "All {techs_all}: {placed} placed, {held} held beside the grid, {ready} ready to assess, {thin} with "
                   "thin evidence, {noev} with no pages yet (most from Tata's list), {off} switched off")],
         before=[("02-technologies.png", "Before: 1,003 'Not assessed', with no reason on any row.")]),
    dict(n="8", title="Every technology classified: commercial availability, domain, TRL, evidence", status="Built",
         source=[("S5", "Tata's own categories: commercial availability, domain and sub-domain, TRL"),
                 ("S1", "Shashank: mature and emerging views"),
                 ("S2", "Akshay's action: represent or filter by TRL"),
                 ("S3", "page 11: the evidence floor counted in independent sources, not pages")],
         what=["Every technology, not only the assessed ones, gets Tata's categories: commercial availability (Research, Pilot, "
               "Early commercial, Mature), domain and sub-domain, and TRL (1–9).",
               "Each answer shows its reason and the pages it rests on, and a person can change it; the change is kept.",
               "The AI also judges how solid the evidence is: how many independent sources, of what kind, and how credible. That, "
               "not a page count, decides when a technology is ready to assess.",
               "The lists are editable in Setup. Sub-domains the AI proposes wait for a person to accept, merge or rename them."],
         why="So Tata sees the radar in its own categories, and early technologies aren't hidden behind a page count.",
         metrics=[("Technologies with availability, domain and TRL", "None", "{classified} of {techs_all}"),
                  ("Agreement with Tata's own list (a sample)", "No classification to compare",
                   "Domain 88%; availability within one step 79%; TRL within one level 9 of 11")],
         before=[("03-technology-agentic-ai.png", "Before: a technology's page showed scores, but no maturity, availability or domain.")]),
    dict(n="9", title="Possible nascent technologies, flagged for a person to confirm", status="Built",
         source=[("S2", "Shashank: nascent technologies shouldn't be hidden below the evidence floor"),
                 ("S1", "Ankur's first test: spot disruptive technologies very early, especially ones Tata doesn't know")],
         what=["The AI flags a technology that looks new and shows a strong early sign (a large funding round, a solid research "
               "result, a first deployment, a spin-out, a new standard), even from a single page, and quotes it.",
               "A person answers yes or no. Yes lets it be assessed despite thin evidence; no is kept until a new page flags it again."],
         why="One deployment can be the signal (the Named Data Networking story), and a page count would hide it.",
         metrics=[("Possibly nascent technologies flagged for a person", "None", "{nascent}, waiting for a yes or no")],
         before=[]),
    dict(n="10", title="Find evidence", status="Built",
         source=[("S1", "Abhishek: open web search beyond the set sources"),
                 ("S3", "page 9: handle unscored measures at Gate 2")],
         what=["A button wherever evidence is missing: on a technology with thin evidence, a measure that couldn't be scored, or a "
               "fabric with no use case.",
               "It searches the web, skips pages already known, reads the new ones like any collected page, and re-runs only the "
               "part that was missing. 'Add a link' beside it takes a page someone already knows.",
               "Only a person starts it; the cost is shown first, and there's a daily limit."],
         why="The set sources can't cover everything, especially for technologies that came in from Tata's list.",
         metrics=[("Ways to look for missing evidence", "None", "Find evidence (web search) and Add a link, wherever evidence is thin")],
         diagram=("find-evidence", "Find evidence searches; Add a link starts from a page a person knows. Both re-run only the gap."),
         before=[("08-assessment-agentic-ai-cloud.png", "Before: an unscored measure said what was missing, with no way to look for it.")]),
    dict(n="11", title="Import a list of emerging technologies, or add one by hand", status="Built",
         source=[("S1", "Abhishek: add technologies by hand, starting from Tata's list"),
                 ("S5", "the 'New Emerging Technologies' tab"),
                 ("S6", "how the upload works: nothing saved until confirmed; Tata's values kept as Tata's")],
         what=["Upload an Excel laid out like Tata's 'New Emerging Technologies' tab (a blank template can be downloaded).",
               "Names are matched to the radar's technologies, including different wordings. Exact names cost nothing; the AI "
               "proposes the rest, and anything unsure stays separate.",
               "Nothing is saved until a person confirms. The row's commercial availability, domain and TRL are kept as Tata's "
               "choice, which a run never overwrites.",
               "One technology can also be added by hand from the technology list."],
         why="Tata already tracks technologies the collection may not have found.",
         metrics=[("Tata's 1,000 rows", "Matched by hand",
                   "Imported on 21 Oct: 88 matched exactly, 27 more matches proposed by the AI and accepted, 885 added")],
         diagram=("import", "Importing Tata's list: a person reviews before anything is saved."),
         before=[]),
    dict(n="12", title="A Point of View that says where the value lands", status="Built",
         source=[("S1", "decided at the review: each Point of View carries the path to scale"),
                 ("S4", "Nick's value levers, blockers and finance"),
                 ("S2", "Hasan: maturity (TRL) stated in the Point of View"),
                 ("S1", "Abhishek: link the sources for market size and other numbers")],
         what=["Up to two use cases per fabric, from the evidence, each with its value lever written as Nick's chain: operating "
               "change → metric → financial effect → evidence needed → next step.",
               "The path to scale in Tri's seven steps (Evaluation → Customer zero → Proof of concept → One customer → Beta → "
               "Roadmap → Investment), recorded by people.",
               "Maturity (TRL) and commercial availability with their reasons, and numbered sources [1], [2] that open the page. "
               "Every number is checked against its source list before the Point of View is saved.",
               "A finance card (contribution margin, payback, break-even), worked out only from figures Tata enters, from Proof of "
               "concept on. The AI never writes a financial figure."],
         why="A Point of View has to say where the value lands for Tata, and every claim has to point to its page.",
         metrics=[("Points of View with numbered sources", "0 of {old_techs}", "{pov_sources} of {povs}, {avg_sources} sources each on average"),
                  ("Use cases with a value lever chain", "None", "Up to 2 per fabric")],
         diagram=("value", "A use case's value as Nick's chain, and the path to scale."),
         before=[("05-pov-agentic-ai.png", "Before: a long text, scores by fabric, no use cases, levers or sources.")]),
    dict(n="13", title="Frontier: the leaders' view, at its own address", status="Built",
         source=[("S1", "decided at the review: Frontier is the 2×2 a business leader sees; Abhishek: separate views for admins "
                        "and users"),
                 ("S1", "Shashank: a table beside the 2×2, fabric filters and zoom"),
                 ("S3", "pages 11–13: the Frontier mock-up")],
         what=["A separate web address for leaders that shows only Frontier: the 2×2 with fabric and stage filters and zoom, a "
               "table beside it, the technologies not placed yet with the reason, and each Point of View, read-only.",
               "It opens on emerging technologies (Research, Pilot, Early commercial); Mature ones are one click away.",
               "Nothing else is reachable from it: no setup, no runs, nothing that costs money.",
               "Not built yet (see 'Not in this round'): publishing editions, the Growth and Defence lenses, the frontier line, "
               "assumptions in force, and leaders recording what they commit."],
         why="Leaders need one clear screen, without the workbench.",
         metrics=[("Addresses", "One admin portal", "Two: Tech Lens for the team, Frontier for leaders"),
                  ("On Frontier (21 Oct)", "No leaders' view",
                   "{fr_placed} placements for {fr_techs} emerging technologies; {fr_mature} at Mature, one click away")],
         diagram=("addresses", "The same portal at two addresses; the leaders' address can only read."),
         before=[]),
]

DECISIONS = [
    "Detection stays the core; each Point of View carries the path to scale; no macro model.",
    "Tata's own categories are used: commercial availability, domain and sub-domain, TRL.",
    "An imported list means 'these are emerging technologies we care about'; each row's own commercial availability is kept.",
    "People's choices are kept and never overwritten by a new run: moves, overrides, hand scores, imports, yes or no on nascent.",
    "The AI never writes a financial figure; contribution, payback and break-even come only from figures Tata enters.",
    "Disruption threat and Serviceable capture are entered by Tata.",
    "Cloud and Security stay separate fabrics for now.",
    "Leaders get their own web address (Frontier).",
    "Assessments use stacked calls (change 5), after trying both ways.",
]

WAITING = [
    "Answer yes or no on the {nascent} possibly nascent technologies.",
    "Accept, merge or rename the {proposed_subdomains} sub-domains the AI proposed.",
    "Settle the measures that need a person, technology by technology, on Gate 2.",
]

ASKS = [
    "Past cases to test against: the FY27 Q1 radar entries and technologies with a known outcome (Named Data Networking, "
    "post-quantum cryptography, Skywave).",
    "Confirmation or correction of each fabric's two value levers and its main blockers (from Nick's note, S4).",
    "Indicative margins, loss rates and delivery costs per fabric, as ranges, and the ROIC proxy when available.",
]

LATER = [
    "Publishing editions, with what moved since the last one (S3, pages 6, 9, 11, 14).",
    "On Frontier: the Growth and Defence lenses, the frontier line, assumptions in force, evidence this edition (S3, page 11), "
    "and leaders recording what they commit (S1; S3, page 13).",
    "Splitting an umbrella term into its emerging part, such as SGP.32 within eSIM (S1).",
    "Tracking inflection points across runs (S1).",
    "Logins, together with a shared environment on Tata's AWS (S3, pages 11, 14, 17).",
    "Calibration against Tata's past cases, and the operating guide for hand-over (S3, pages 14, 17).",
    "Licensed sources through data providers or MCPs (S1).",
    "Combining Cloud and Security (on hold). Collection on a timer will not be built.",
]


def t(text: str) -> str:
    """Text with the numbers filled in."""
    return text.format(**N)


def shade(cell, hex_fill):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tc.append(shd)


def table(doc, header, rows, widths):
    tb = doc.add_table(rows=1, cols=len(header))
    tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = tb.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "041A4A")
    for row in rows:
        cells = tb.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = t(str(v))
    tb.autofit = False
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
    tb._tbl.tblPr.append(layout)
    for i, w in enumerate(widths):
        tb.columns[i].width = Inches(w)
        for cell in tb.columns[i].cells:
            cell.width = Inches(w)
    return tb


def image(doc, path: Path, caption: str):
    if not path.exists():
        doc.add_paragraph(f"[{path.name} not found]"); return
    doc.add_picture(str(path), width=Inches(6.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(caption); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: r.italic = True; r.font.size = Pt(9); r.font.color.rgb = MUTED


def label(doc, text: str):
    doc.add_paragraph().add_run(text).bold = True


def main():
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for s in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        doc.styles[s].font.color.rgb = NAVY
    sec = doc.sections[0]
    for side in ("left_margin", "right_margin"): setattr(sec, side, Inches(1))

    doc.add_paragraph("Tech Radar: changes after the midpoint review", style="Title")
    p = doc.add_paragraph("Tata Communications × Wharton Mack Institute · Emerging Technology Radar"); p.runs[0].font.color.rgb = MUTED
    p = doc.add_paragraph(VERSION); p.runs[0].bold = True

    doc.add_heading("About this document", level=1)
    doc.add_paragraph("Every change we made after the midpoint review of 13 October, built from the review itself, the emails "
                      "and notes that followed, and our own working sessions on them. Each change has:")
    for b in ["Where it came from: the email or document (S1–S6, listed below) and the point, page or person in it.",
              "What changed, in plain words, and why.",
              "Before → after, in numbers, and the screens before and after. A workflow diagram where it helps.",
              "Numbers come from the radar itself: 'before' from the run of 6 October, 'after' from the full run of 21 October."]:
        doc.add_paragraph(b, style="List Bullet")

    doc.add_heading("Where the changes came from", level=1)
    table(doc, ["", "Email or document", "What it is"], SOURCES, [0.4, 2.6, 3.5])

    doc.add_heading("Summary", level=1)
    table(doc, ["#", "Change", "Came from", "Status"],
          [(i["n"], i["title"], ", ".join(dict.fromkeys(s for s, _ in i["source"])), i["status"]) for i in ITEMS],
          [0.4, 3.6, 1.6, 0.9])

    doc.add_heading("How the radar works now", level=1)
    doc.add_paragraph("The radar's steps after these changes. Filled boxes are new or changed; the amber lanes are where a person decides.")
    image(doc, DIAGRAMS / "process.png", "The process, from collecting pages to the leaders' view.")

    doc.add_heading("The changes, one by one", level=1)
    for it in ITEMS:
        doc.add_heading(f"{it['n']}. {it['title']}", level=2)
        p = doc.add_paragraph(); r = p.add_run(f"Status: {it['status']}"); r.font.color.rgb = MUTED; r.font.size = Pt(9.5)
        label(doc, "Where it came from")
        for code, where in it["source"]:
            p = doc.add_paragraph(style="List Bullet"); p.add_run(f"{code}: ").bold = True; p.add_run(where)
        label(doc, "What changed")
        for w in it["what"]: doc.add_paragraph(t(w), style="List Bullet")
        p = doc.add_paragraph(); p.add_run("Why: ").bold = True; p.add_run(t(it["why"]))
        label(doc, "Before → after")
        table(doc, ["", "Before", "After"], it["metrics"], [2.0, 2.2, 2.3])
        if it.get("diagram"):
            label(doc, "How it works")
            image(doc, DIAGRAMS / f"{it['diagram'][0]}.png", it["diagram"][1])
        label(doc, "Screens")
        for name, caption in it["before"]:
            image(doc, BEFORE / name, caption)
        if not it["before"]:
            doc.add_paragraph("Before: this didn't exist.")
        built = [AFTER_BUILT / n for n in AFTER_CAPTIONS if n.startswith(f"item-{it['n']}-")]
        for b in [b for b in built if b.exists()]:
            image(doc, b, AFTER_CAPTIONS[b.name])

    doc.add_heading("Decisions taken since the review", level=1)
    for d in DECISIONS: doc.add_paragraph(d, style="List Bullet")

    doc.add_heading("Waiting for Tata's team, in the portal", level=1)
    for w in WAITING: doc.add_paragraph(t(w), style="List Number")

    doc.add_heading("What we need from Tata", level=1)
    for a in ASKS: doc.add_paragraph(a, style="List Number")

    doc.add_heading("Not in this round", level=1)
    for l in LATER: doc.add_paragraph(l, style="List Bullet")

    doc.add_heading("Change history", level=1)
    table(doc, ["Version", "Date", "What changed"], HISTORY, [0.8, 1.2, 4.5])

    doc.save(OUT)
    print("wrote", OUT, OUT.stat().st_size // 1024, "KB")


if __name__ == "__main__":
    main()
