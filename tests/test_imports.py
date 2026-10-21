"""Import: the header sheet found, a wrong file refused with its reason,
exact, proposed and new groups, a rejected match comes in new, confirm
applies aliases and choices, importing again changes nothing, and imported
names stay in the organize step's register. Then the fixes of 21 Oct: code
keeps a narrower name apart, the name in brackets proposes a match for free,
the review shows what a match would join, several rows naming one technology
leave their values on the record, a sub-domain not in our lists waits for a
person, reopen undoes a confirm for some rows (also one recorded the old
way), and adding by hand shows its cost and takes values."""

import io
import json

import openpyxl
from fastapi.testclient import TestClient

from app.main import app
from app.services import classify_effective, decisions, imports
from app.services.timezone import now_utc
from agent.resolver.resolver import known_register
from factories import tech
from stacked_fakes import load_seeded

client = TestClient(app)
HEAD = ["Technology Name", "Domain", "Sub-domain", "Description", "TRL", "Commercial Availability"]
TZ = "America/Chicago"
FOLDS = {"llm agents": "agentic ai", "mist computing": "edge computing", "edge cloud": "edge computing",
         "telco edge computing": "edge computing", "computer vision for robotics": "computer vision",
         "6g cellular communications": "6g"}


def book(rows, other_first=True) -> bytes:
    wb = openpyxl.Workbook()
    if other_first:
        wb.active.title = "Technologies Tracked"
        wb.active.append([None, "Technology"])
        ws = wb.create_sheet("New Emerging Technologies")
    else:
        ws = wb.active
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class Resolver:
    """Folds the names in FOLDS that it is sent; everything else stays apart."""
    def ask(self, prompt, system=None, **kw):
        raw = json.loads(prompt.rsplit("Raw technology names:\n", 1)[1])
        return json.dumps({"groups": [{"name": FOLDS[n], "members": [n]} for n in raw if n in FOLDS]})


def upload(content, name="list.xlsx"):
    return client.post("/api/imports", files={"file": (name, content)})


def _world(db, monkeypatch):
    load_seeded(db, "prompts", "availability", "domains", "subdomains", "trl_levels")
    tech(db, "agentic-ai", "agentic ai")
    tech(db, "edge-ai", "edge ai")
    from app.services.llm import llm as llm_module
    monkeypatch.setattr(llm_module.LLM, "initialize", classmethod(lambda cls, *a, **k: Resolver()))


def _moves(db, tech_id: str, active=True) -> dict:
    return {d["field"]: d["value"] for d in db["decisions"].find({"tech_id": tech_id, "kind": "move", "active": active})}


def _decide(imp_id: str, rows: list[int], decision: str) -> dict:
    return client.put(f"/api/imports/{imp_id}/items", json={"rows": rows, "decision": decision}).json()


def test_a_wrong_file_is_refused_with_why(db, monkeypatch):
    _world(db, monkeypatch)
    assert "only .xlsx" in upload(b"a,b", "list.csv").json()["detail"]
    wb = openpyxl.Workbook(); wb.active.append(["Name"]); out = io.BytesIO(); wb.save(out)
    assert upload(out.getvalue()).json()["detail"] == "no 'Technology Name' column in any sheet"


def test_groups_review_confirm_and_again(db, monkeypatch):
    _world(db, monkeypatch)
    rows = [["Edge AI", "AI", "Edge AI", "d", "6", "Early Commercial"],
            ["LLM agents", "AI", "AI Agents", "d", "", "Pilot"],
            ["Brain chips", "AI", "Neuromorphic Computing", "d", "3", "Research"],
            ["edge ai", "AI", "", "", "", ""]]                              # a repeat
    imp = upload(book(rows)).json()
    assert imp["sheet"] == "New Emerging Technologies" and imp["rows"] == 4 and len(imp["items"]) == 3
    assert {i["name"]: i["group"] for i in imp["items"]} == {"Edge AI": "matched", "LLM agents": "unmatched",
                                                             "Brain chips": "unmatched"}
    r = client.post(f"/api/imports/{imp['_id']}/confirm")
    assert r.status_code == 409 and "match the names first" in r.json()["detail"]
    imp = client.post(f"/api/imports/{imp['_id']}/match").json()
    groups = {i["name"]: (i["group"], i["target_tech_id"]) for i in imp["items"]}
    assert groups["LLM agents"] == ("proposed", "agentic-ai") and groups["Brain chips"][0] == "new"
    assert client.post(f"/api/imports/{imp['_id']}/confirm").status_code == 409  # a proposal waits for a person
    row = next(i["row"] for i in imp["items"] if i["name"] == "LLM agents")
    _decide(imp["_id"], [row], "accept")
    assert client.post(f"/api/imports/{imp['_id']}/confirm").json() == {"joined": 2, "added": 1, "shared": 0}
    agentic = db["technologies"].find_one({"_id": "agentic-ai"})
    assert "llm agents" in agentic["aliases"] and agentic["imported"]["file"] == "list.xlsx"
    assert _moves(db, "agentic-ai") == {"availability": "pilot", "domain": "ai", "subdomain": "ai--ai_agents"}
    brain = db["technologies"].find_one({"_id": "brain-chips"})
    assert brain["origin"] == "imported" and brain["page_count"] == 0
    e = classify_effective.effective(brain, decisions.active(["brain-chips"]), classify_effective.names())
    assert (e["availability"]["id"], e["availability"]["source"], e["trl"]["level"]) == ("research", "person", 3)
    assert "brain chips" in known_register(1)
    before = db["decisions"].count_documents({})
    again = upload(book(rows)).json()
    assert all(i["group"] == "matched" for i in again["items"])
    client.post(f"/api/imports/{again['_id']}/confirm")
    assert db["decisions"].count_documents({}) == before and db["technologies"].count_documents({}) == 3


def test_a_rejected_match_comes_in_new(db, monkeypatch):
    _world(db, monkeypatch)
    imp = upload(book([["LLM agents", "AI", "", "", "", "Pilot"]], other_first=False)).json()
    imp = client.post(f"/api/imports/{imp['_id']}/match").json()
    _decide(imp["_id"], [imp["items"][0]["row"]], "reject")
    assert client.post(f"/api/imports/{imp['_id']}/confirm").json() == {"joined": 0, "added": 1, "shared": 0}
    assert db["technologies"].find_one({"_id": "llm-agents"})["origin"] == "imported"


def test_a_narrower_name_is_never_the_broader_one():
    edge = ["edge computing", "Edge computing"]
    for name in ("Telco Edge Computing", "Tactical Edge Computing", "Decentralized Edge Computing"):
        assert imports.narrower(name, edge), name
    assert imports.narrower("Computer Vision for Robotics", ["computer vision"])
    assert imports.narrower("AI in Education", ["artificial intelligence"])
    assert imports.narrower("Federated Learning at Edge", ["federated learning"])
    assert imports.narrower("Generative AI", ["artificial intelligence", "ai"])
    # the same thing in other words, or worded at more length, stays the skill's call
    assert not imports.narrower("6G Cellular Communications", ["6g"])
    assert not imports.narrower("Co-Packaged Optics (CPO)", ["co-packaged optics"])
    assert not imports.narrower("DNS Security Extensions (DNSSEC)", ["dnssec"])
    assert not imports.narrower("Network-on-Chip (NoC)", ["network on chip"])
    assert not imports.narrower("Mist Computing", edge)


def test_matching_keeps_narrower_names_apart_and_reads_brackets(db, monkeypatch):
    _world(db, monkeypatch)
    for t_id, name in (("edge-computing", "edge computing"), ("computer-vision", "computer vision"), ("6g", "6g"),
                       ("dnssec", "dnssec"), ("dns", "domain name system (dns)")):
        tech(db, t_id, name)
    db["technologies"].update_one({"_id": "edge-computing"}, {"$set": {"page_count": 15}})
    imp = upload(book([["Telco Edge Computing", "Edge Computing", "", "", "", ""],
                       ["Computer Vision for Robotics", "", "", "", "", ""],
                       ["6G Cellular Communications", "", "", "", "", ""],
                       ["Mist Computing", "", "", "", "", ""],
                       ["DNS Security Extensions (DNSSEC)", "", "", "", "", ""],
                       ["Edge AI (IoT)", "", "", "", "", ""]])).json()
    by = {i["name"]: i for i in imp["items"]}
    # the name in its brackets, or without them, proposes a match at no cost
    assert (by["DNS Security Extensions (DNSSEC)"]["group"], by["DNS Security Extensions (DNSSEC)"]["target_tech_id"]) == \
        ("proposed", "dnssec") and "DNSSEC" in by["DNS Security Extensions (DNSSEC)"]["why"]
    assert (by["Edge AI (IoT)"]["group"], by["Edge AI (IoT)"]["target_tech_id"]) == ("proposed", "edge-ai")
    assert client.get(f"/api/imports/{imp['_id']}/match/preview").json()["names"] == 4
    imp = client.post(f"/api/imports/{imp['_id']}/match").json()
    by = {i["name"]: i for i in imp["items"]}
    assert by["Telco Edge Computing"]["group"] == "new"
    assert by["Telco Edge Computing"]["kept_apart"] == {"tech_id": "edge-computing", "name": "edge computing"}
    assert by["Computer Vision for Robotics"]["group"] == "new"
    assert (by["6G Cellular Communications"]["group"], by["Mist Computing"]["group"]) == ("proposed", "proposed")
    # the review has what's needed to judge a match: the technology's pages and classification
    edge = imp["targets"]["edge-computing"]
    assert edge["page_count"] == 15 and edge["availability"]["name"] == "Unclassified" and edge["trl"]["level"] is None


def test_several_rows_naming_one_technology_leave_their_values_on_the_record(db, monkeypatch):
    _world(db, monkeypatch)
    tech(db, "edge-computing", "edge computing")
    rows = [["Edge Computing", "Edge Computing", "Distributed Computing", "d", "", "Early Commercial"],
            ["Mist Computing", "Edge Computing", "Edge Data Processing", "d", "6", "Research"],
            ["Edge Cloud", "Cloud", "Edge-Cloud Convergence", "d", "", "Early Commercial"],
            ["LLM agents", "AI", "AI Agents", "d", "", "Pilot"]]
    imp = upload(book(rows)).json()
    imp = client.post(f"/api/imports/{imp['_id']}/match").json()
    _decide(imp["_id"], [i["row"] for i in imp["items"] if i["group"] == "proposed"], "accept")
    assert client.post(f"/api/imports/{imp['_id']}/confirm").json() == {"joined": 4, "added": 0, "shared": 3}
    assert _moves(db, "edge-computing") == {}                                  # no row's values won
    rec = db["technologies"].find_one({"_id": "edge-computing"})["imported"]
    assert rec["rows"] == ["Edge Computing", "Mist Computing", "Edge Cloud"] and "row" not in rec
    assert _moves(db, "agentic-ai")["availability"] == "pilot"                 # one row alone: its values
    before = db["decisions"].count_documents({})
    again = upload(book(rows)).json()                                          # every name is known now
    assert {i["group"] for i in again["items"]} == {"matched"}
    assert client.post(f"/api/imports/{again['_id']}/confirm").json()["shared"] == 3
    assert db["decisions"].count_documents({}) == before


def test_importing_again_keeps_a_persons_later_change(db, monkeypatch):
    _world(db, monkeypatch)
    rows = [["Brain chips", "AI", "", "d", "3", "Research"]]
    imp = upload(book(rows)).json()
    client.post(f"/api/imports/{imp['_id']}/match")
    client.post(f"/api/imports/{imp['_id']}/confirm")
    decisions.create("move", "brain-chips", field="availability", value="pilot", reason="seen in a trial", tz=TZ)
    again = upload(book(rows)).json()
    client.post(f"/api/imports/{again['_id']}/confirm")
    assert _moves(db, "brain-chips")["availability"] == "pilot"


def test_a_subdomain_not_in_our_lists_waits_for_a_person(db, monkeypatch):
    _world(db, monkeypatch)
    imp = upload(book([["Brain chips", "AI", "Brainware", "d", "", ""],
                       ["Mind chips", "Neuro", "Wetware", "d", "", ""]])).json()
    client.post(f"/api/imports/{imp['_id']}/match")
    imp = client.get(f"/api/imports/{imp['_id']}").json()
    sub = next(u for u in imp["unknown_values"] if u["field"] == "subdomain")
    assert (sub["value"], sub["keep"], imp["counts"]["values_waiting"]) == ("Brainware", None, 1)
    r = client.post(f"/api/imports/{imp['_id']}/confirm")
    assert r.status_code == 409 and "sub-domain" in r.json()["detail"] and "1 still waits" in r.json()["detail"]
    # mapping the unknown domain brings its sub-domain in, to choose for too
    imp = client.put(f"/api/imports/{imp['_id']}/values", json=[{"field": "domain", "value": "Neuro", "map_to": "ai"}]).json()
    assert imp["counts"]["values_waiting"] == 2
    imp = client.put(f"/api/imports/{imp['_id']}/values", json=[
        {"field": "subdomain", "value": "Brainware", "domain_id": "ai", "keep": True},
        {"field": "subdomain", "value": "Wetware", "domain_id": "ai", "map_to": None, "keep": False}]).json()
    assert imp["counts"]["values_waiting"] == 0 and imp["counts"]["values_to_map"] == 1
    assert client.post(f"/api/imports/{imp['_id']}/confirm").status_code == 200
    assert _moves(db, "brain-chips")["subdomain"] == "new:Brainware"
    assert _moves(db, "mind-chips") == {"domain": "ai"}


def _old_confirm(db, import_id: str) -> None:
    """Confirm as it worked until 21 Oct, as the live import was confirmed:
    every joining row's values recorded in turn, last row winning, without
    naming the row."""
    imp = db["imports"].find_one({"_id": import_id})
    for it in imp["items"]:
        record = {"import_id": import_id, "file": imp["file"], "at": now_utc(), "tz": TZ,
                  "row": {**it["values"], "name": it["name"]}}
        if it["group"] == "matched" or it["decision"] == "accept":
            tech_id = it["target_tech_id"]
            db["technologies"].update_one({"_id": tech_id}, {"$addToSet": {"aliases": it["name"].lower()},
                                                             "$set": {"imported": record}})
        else:
            tech_id = imports.new_technology(it["name"], "imported", {"imported": record})
        active = {d["field"]: d["value"] for d in decisions.active([tech_id], ["move"])}
        for field, value in imports.resolved_values(it["values"], imp["unknown_values"]).items():
            if active.get(field) != value:
                decisions.create("move", tech_id, field=field, value=value, reason=f"Imported from {imp['file']}", tz=TZ)
    db["imports"].update_one({"_id": import_id}, {"$set": {"status": "confirmed", "confirmed_at": now_utc(),
                                                           "result": {"joined": 4, "added": 1}}})


REOPEN_ROWS = [["Edge Computing", "Edge Computing", "Distributed Computing", "d", "", "Early Commercial"],
               ["Mist Computing", "Edge Computing", "Edge Data Processing", "d", "6", "Research"],
               ["Edge Cloud", "Cloud", "Edge-Cloud Convergence", "d", "", "Early Commercial"],
               ["LLM agents", "AI", "AI Agents", "d", "", "Pilot"],
               ["Brain chips", "AI", "Neuromorphic Computing", "d", "3", "Research"]]


def _reopen_world(db, monkeypatch) -> dict:
    _world(db, monkeypatch)
    tech(db, "edge-computing", "edge computing")
    decisions.create("move", "edge-computing", field="trl", value=5, reason="read the pages", tz=TZ)  # its own
    imp = upload(book(REOPEN_ROWS)).json()
    imp = client.post(f"/api/imports/{imp['_id']}/match").json()
    _decide(imp["_id"], [i["row"] for i in imp["items"] if i["group"] == "proposed"], "accept")
    return {i["name"]: i["row"] for i in imp["items"]} | {"_id": imp["_id"]}


def test_reopen_undoes_an_old_style_confirm_and_confirming_again_is_safe(db, monkeypatch):
    rows = _reopen_world(db, monkeypatch)
    imp_id = rows["_id"]
    _old_confirm(db, imp_id)
    assert _moves(db, "edge-computing") == {"domain": "cloud", "subdomain": "cloud--edge_cloud_convergence",
                                            "availability": "early_commercial", "trl": 6}   # the last row won
    brain = {d["_id"] for d in db["decisions"].find({"tech_id": "brain-chips"})}
    assert client.post(f"/api/imports/{imp_id}/reopen", json={"rows": [rows["Brain chips"]]}).status_code == 409

    r = client.post(f"/api/imports/{imp_id}/reopen").json()                 # every accepted proposed match
    assert r == {"reopened": 3, "values_withdrawn": 7, "restored": 1}
    edge = db["technologies"].find_one({"_id": "edge-computing"})
    assert edge["aliases"] == ["edge computing"]
    assert _moves(db, "edge-computing") == {"trl": 5}                       # its own value is back
    assert edge["imported"]["row"]["name"] == "Edge Computing"              # the row that still names it
    agentic = db["technologies"].find_one({"_id": "agentic-ai"})
    assert "llm agents" not in agentic["aliases"] and "imported" not in agentic and _moves(db, "agentic-ai") == {}
    imp = client.get(f"/api/imports/{imp_id}").json()
    assert imp["status"] == "ready" and imp["counts"]["accepted"] == 0 and imp["counts"]["proposed"] == 3
    assert client.post(f"/api/imports/{imp_id}/confirm").status_code == 409    # back to review
    assert client.post(f"/api/imports/{imp_id}/reopen").status_code == 409     # not confirmed now

    _decide(imp_id, [rows["Mist Computing"]], "reject")
    _decide(imp_id, [rows["Edge Cloud"], rows["LLM agents"]], "accept")
    assert client.post(f"/api/imports/{imp_id}/confirm").json() == {"joined": 3, "added": 1, "shared": 2}
    assert _moves(db, "edge-computing") == {"trl": 5}                       # two rows name it: neither's values
    assert _moves(db, "mist-computing") == {"domain": "edge_computing", "subdomain": "edge_computing--edge_data_processing",
                                            "trl": 6, "availability": "research"}
    assert _moves(db, "agentic-ai")["availability"] == "pilot"
    assert {d["_id"] for d in db["decisions"].find({"tech_id": "brain-chips"})} == brain   # untouched: nothing new
    assert db["technologies"].count_documents({"origin": "imported"}) == 2


def test_confirm_reopen_confirm_again_changes_nothing_for_the_other_rows(db, monkeypatch):
    rows = _reopen_world(db, monkeypatch)
    imp_id = rows["_id"]
    assert client.post(f"/api/imports/{imp_id}/confirm").json() == {"joined": 4, "added": 1, "shared": 3}
    assert _moves(db, "edge-computing") == {"trl": 5}
    state = {t: _moves(db, t) for t in ("agentic-ai", "brain-chips", "edge-computing")}
    count = db["decisions"].count_documents({})

    r = client.post(f"/api/imports/{imp_id}/reopen", json={"rows": [rows["LLM agents"]]}).json()
    assert r == {"reopened": 1, "values_withdrawn": 3, "restored": 0}
    assert _moves(db, "agentic-ai") == {} and _moves(db, "brain-chips") == state["brain-chips"]
    _decide(imp_id, [rows["LLM agents"]], "accept")
    client.post(f"/api/imports/{imp_id}/confirm")
    assert {t: _moves(db, t) for t in state} == state                       # the same as the first confirm
    assert db["decisions"].count_documents({}) == count + 3                 # only the reopened row's values again


def test_add_one_by_hand(db, monkeypatch):
    _world(db, monkeypatch)
    tech(db, "dnssec", "dnssec")
    p = client.get("/api/technologies/add/preview", params={"name": "Edge AI"}).json()
    assert p["outcome"] == "matched" and p["cost"] is None and p["target"]["tech_id"] == "edge-ai"
    p = client.get("/api/technologies/add/preview", params={"name": "DNS Security Extensions (DNSSEC)"}).json()
    assert p["outcome"] == "proposed" and p["cost"] is None and p["target"]["tech_id"] == "dnssec"
    p = client.get("/api/technologies/add/preview", params={"name": "LLM agents"}).json()
    assert p["outcome"] == "unmatched" and p["cost"]["calls"] == 1 and p["cost"]["register"] == 3
    assert p["cost"]["tokens_about"] > 300                                  # the skill's text and the register

    assert client.post("/api/technologies", json={"name": "Edge AI"}).json()["outcome"] == "matched"
    p = client.post("/api/technologies", json={"name": "LLM agents"}).json()
    assert p["outcome"] == "proposed" and p["proposed"]["tech_id"] == "agentic-ai" and "page_count" in p["proposed"]
    r = client.post("/api/technologies", json={"name": "LLM agents", "confirm": "new", "availability": "research",
                                              "domain": "ai", "subdomain": "ai--ai_agents", "trl": 4}).json()
    assert r["outcome"] == "added" and r["values_set"] == 4
    added = db["technologies"].find_one({"_id": "llm-agents"})
    assert added["origin"] == "added"
    e = classify_effective.effective(added, decisions.active(["llm-agents"]), classify_effective.names())
    assert (e["availability"]["id"], e["availability"]["source"], e["availability"]["emerging"]) == ("research", "person", True)
    bad = client.post("/api/technologies", json={"name": "X", "domain": "ai", "subdomain": "cloud--cloud_compute"})
    assert bad.status_code == 422 and "isn't a sub-domain of 'ai'" in bad.json()["detail"]
    assert client.post("/api/technologies", json={"name": "Y", "availability": "soon"}).status_code == 422
